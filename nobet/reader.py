"""Genel belge okuyucu: PDF, Excel ve CSV dosyalarından tablo + metin çıkarır.

Her belge "sayfa" listesine dönüşür. Bir sayfa: {"text": str, "tables": [[[hücre, ...], ...], ...]}
Excel'de her çalışma sayfası bir sayfa olarak döner.
"""
import csv
import os


class ReadError(Exception):
    pass


def read_document(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        return _read_pdf(path)
    if ext in (".xlsx", ".xlsm"):
        return _read_xlsx(path)
    if ext == ".csv":
        return _read_csv(path)
    raise ReadError(f"Desteklenmeyen dosya türü: {ext} (PDF, XLSX veya CSV seçin)")


def _cell(v):
    if v is None:
        return ""
    return str(v).replace("\r", "").strip()


def _read_pdf(path):
    try:
        import pdfplumber
    except ImportError as e:  # pragma: no cover
        raise ReadError("pdfplumber kurulu değil") from e
    pages = []
    try:
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                text = page.extract_text() or ""
                tables = []
                # Önce çizgili tablo, bulunamazsa metin hizasına göre tablo dene.
                for settings in (None, {"vertical_strategy": "text", "horizontal_strategy": "lines"},
                                 {"vertical_strategy": "text", "horizontal_strategy": "text"}):
                    try:
                        found = page.find_tables(settings) if settings else page.find_tables()
                    except Exception:
                        found = []
                    for tf in found:
                        # Küçük yazılı iki satır birbirine karışmasın (ör. "MAT. ve BİL. / UYG.")
                        data = tf.extract(y_tolerance=1, x_tolerance=1.5)
                        if not data or len(data) < 2 or max(len(r) for r in data) < 3:
                            continue
                        data = [[_cell_raw(c) for c in row] for row in data]
                        if settings is None:
                            _merge_box_spans(data, tf, page)
                        tables.append(data)
                    if tables:
                        break
                notes = []
                try:
                    notes = _audit_classes(page)
                except Exception:  # noqa: BLE001 — denetim başarısızsa okuma yine de sürer
                    notes = []
                pages.append({"text": text, "tables": tables, "notes": notes})
    except ReadError:
        raise
    except Exception as e:
        raise ReadError(f"PDF okunamadı: {e}") from e
    if not any(p["text"].strip() or p["tables"] for p in pages):
        raise ReadError("PDF'te okunabilir metin yok. Belge taranmış (resim) olabilir; "
                        "e-Okul/MEBBİS'ten PDF veya Excel olarak yeniden indirin.")
    return pages


def _audit_classes(page):
    """Okumayı belgeyle karşılaştırır: her öğretmen satırında PDF'te yazan sınıf adları,
    okunan hücrelerde de bulunmalı. Bulunmayan sınıf varsa uyarı döner."""
    import re
    notes = []
    class_re = re.compile(r"(?<![\w/])(\d{1,2}\s*[-/.]?\s*[A-ZÇĞİÖŞÜ])(?![\wçğıöşü])")
    for tf in page.find_tables():
        data = tf.extract(y_tolerance=1, x_tolerance=1.5)
        for ri, row in enumerate(tf.rows):
            if ri >= len(data) or not data[ri] or not data[ri][0]:
                continue
            name = str(data[ri][0]).replace("\n", " ").strip()
            last = [str(c or "").strip() for c in data[ri]][-1]
            if not last.isdigit():
                continue                                  # öğretmen satırı değil
            cells = [c for c in row.cells if c]
            if len(cells) < 3:
                continue
            x0, x1 = cells[0][2], cells[-1][0]
            words = page.within_bbox(row.bbox).extract_words(y_tolerance=1, x_tolerance=1.5)
            body = " ".join(w["text"] for w in words if w["x0"] >= x0 - 1 and w["x1"] <= x1 + 1)
            in_pdf = {m.group(1).replace(" ", "") for m in class_re.finditer(body)}
            # Yan yana yapışık sınıflar ("5A5G") ayrı sayılır
            for w in words:
                for m in re.finditer(r"\d{1,2}\s*[/.-]?\s*[A-ZÇĞİÖŞÜ](?![a-zçğıöşü])", w["text"]):
                    in_pdf.add(m.group(0).replace(" ", ""))
            seen = " ".join(str(c or "") for c in data[ri][1:-1])
            in_cells = {m.group(0).replace(" ", "") for m in
                        re.finditer(r"\d{1,2}\s*[/.-]?\s*[A-ZÇĞİÖŞÜ](?![a-zçğıöşü])", seen)}
            lost = sorted(in_pdf - in_cells)
            if lost:
                notes.append(f"⚠ {name}: belgede görünen sınıf(lar) okunamadı: {', '.join(lost)}")
    return notes


def _merge_box_spans(data, tf, page):
    """Birden fazla sütunu kaplayan ders kutularını (blok ders) tüm sütunlara yayar.

    Bazı programlar (ör. e-Okul çarşaf listesi) her dersi ayrı bir dikdörtgen olarak çizer;
    2-3 saatlik blok derste kutu birkaç sütunu kaplar ama yazı tek bir sütuna, bazen de
    iki sütuna bölünmüş olarak düşer. Kutunun kapladığı her hücreye aynı metni yazarız.
    """
    for ri, row in enumerate(tf.rows):
        if ri >= len(data):
            break
        top, bottom = row.bbox[1], row.bbox[3]
        height = bottom - top
        centers = [None if c is None else (c[0] + c[2]) / 2 for c in row.cells]
        limit = max(2, len(centers) // 3)
        for r in page.rects:
            h = r["bottom"] - r["top"]
            if r["top"] < top - 1.5 or r["bottom"] > bottom + 1.5 or h < height * 0.5:
                continue
            cols = [j for j, cx in enumerate(centers)
                    if cx is not None and j < len(data[ri]) and r["x0"] - 0.5 <= cx <= r["x1"] + 0.5]
            if not 2 <= len(cols) <= limit:
                continue
            texts = []
            for j in cols:
                v = data[ri][j]
                if v and v not in texts:
                    texts.append(v)
            if not texts:
                continue
            # Yazı iki hücrenin sınırına düşüp bölünmüş olabilir; kutunun içinden yeniden oku.
            try:
                box = page.within_bbox((r["x0"] - 0.5, r["top"] - 0.5, r["x1"] + 0.5,
                                        r["bottom"] + 0.5))
                merged = (box.extract_text(y_tolerance=1, x_tolerance=1.5) or "").strip()
            except Exception:
                merged = ""
            merged = merged or " ".join(texts)
            for j in cols:
                data[ri][j] = merged


def _cell_raw(v):
    # PDF'te birleştirilmiş hücreler None döner; bunu "" yerine None olarak korumak
    # başlıkları sağa doğru doldurmak için gerekli.
    if v is None:
        return None
    return str(v).replace("\r", "").strip()


def _read_xlsx(path):
    try:
        import openpyxl
    except ImportError as e:  # pragma: no cover
        raise ReadError("openpyxl kurulu değil") from e
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    except Exception as e:
        raise ReadError(f"Excel dosyası açılamadı: {e}") from e
    pages = []
    for ws in wb.worksheets:
        rows = []
        merged_tail = set()
        for rng in ws.merged_cells.ranges:
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    if (r, c) != (rng.min_row, rng.min_col):
                        merged_tail.add((r, c))
        for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            cells = []
            for c_idx, v in enumerate(row, start=1):
                cells.append(None if (r_idx, c_idx) in merged_tail else _cell(v))
            if any(cells):
                rows.append(cells)
        text = "\n".join(" ".join(c for c in r if c) for r in rows)
        pages.append({"text": text, "tables": [rows] if rows else []})
    return pages


def _read_csv(path):
    for enc in ("utf-8-sig", "cp1254", "latin-1"):
        try:
            with open(path, newline="", encoding=enc) as f:
                sample = f.read(4096)
                f.seek(0)
                try:
                    dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
                except csv.Error:
                    dialect = csv.excel
                rows = [[_cell(c) for c in r] for r in csv.reader(f, dialect)]
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ReadError("CSV kodlaması çözülemedi")
    rows = [r for r in rows if any(r)]
    text = "\n".join(" ".join(c for c in r if c) for r in rows)
    return [{"text": text, "tables": [rows] if rows else []}]
