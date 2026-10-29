"""Excel ve yazdırma (HTML) çıktıları."""
import datetime as dt
import html

from .names import DAYS, lesson_label, normalize, split_info
from .scheduler import KIND_LABELS


def date_tr(date):
    d = dt.date.fromisoformat(date)
    months = ["Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran", "Temmuz", "Ağustos",
              "Eylül", "Ekim", "Kasım", "Aralık"]
    day = DAYS[d.weekday()] if d.weekday() < 5 else ("Cumartesi", "Pazar")[d.weekday() - 5]
    return f"{d.day} {months[d.month - 1]} {d.year} {day}"


def month_tr(month):
    y, m = month.split("-")
    months = ["Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran", "Temmuz", "Ağustos",
              "Eylül", "Ekim", "Kasım", "Aralık"]
    return f"{months[int(m) - 1]} {y}"


def day_html(plan, info=None, has_logo=False):
    """Yazdırılacak resmî görevlendirme çizelgesi (A4 dikey)."""
    e = html.escape
    info = info if isinstance(info, dict) else {"school": info or ""}
    school = info.get("school", "")
    city = info.get("school_city", "")
    tables = []
    groups = {}
    for a in plan.assignments:
        groups.setdefault(a.absent, []).append(a)
    for absent, items in groups.items():
        rows = []
        reason = plan.absences.get(absent, "")
        rows.append(
            f"<tr style='background:#e2e8f0'><td colspan=5><b>{e(absent)}</b>"
            f" &nbsp;—&nbsp; Mazeret: <b>{e(reason)}</b> &nbsp;({len(items)} ders)</td></tr>")
        for a in items:
            color = {"bos": "#fff6db", "yok": "#fde7e7", "joker": "#f1e6ff"}.get(a.kind, "#ffffff")
            sinif, ders = split_info(a.info)
            ders = lesson_label(ders)
            note = e(a.note.replace("⚠ ", "")) if a.note else ""
            rows.append(
                f"<tr style='background:{color}'><td align=center width=9%><b>{a.hour}. ders</b></td>"
                f"<td align=center width=12%>{e(sinif) or '—'}</td><td width=22%>{e(ders or a.info)}</td>"
                f"<td><b>{e(a.assigned or '— atanmadı —')}</b>"
                + (f"<br><span style='color:#92400e; font-size:7pt'>{note}</span>" if note else "") +
                "</td><td width=18%></td></tr>")
        # Her gelmeyen öğretmenin tablosu ayrı; aralarında bir satır boşluk olur
        tables.append(
            "<table border=1 cellspacing=0 cellpadding=5 width=100% "
            "style='border-collapse:collapse'>"
            "<tr style='background:#cbd5e1'><th>Saat</th><th>Sınıf</th><th>Ders</th>"
            "<th>Derse Girecek Öğretmen</th><th>İmza</th></tr>" + "".join(rows) + "</table>")
    duty = ", ".join(f"{e(t)}" + (f" ({e(p)})" if p else "") for t, p in plan.duty) or "—"
    logo = "<img src='logo' height=70>" if has_logo else ""
    head = "<br>".join(x for x in ("T.C.", e(city), f"<b>{e(school)}</b>") if x)
    sign = []
    for title, key in (("Müdür Yardımcısı", "vice_principal"), ("Okul Müdürü", "principal")):
        name = (info.get(key, "") or "").replace(";", ",").split(",")[0].strip()
        sign.append(f"<td align=center width=40%>{e(name) or '&nbsp;'}<br>"
                    f"<b>{title}</b><br><br>İmza</td>")
    return f"""
<html><body style="font-family:'Segoe UI', Arial; font-size:9pt; color:#111">
<table width=100% cellspacing=0 cellpadding=0><tr>
<td width=15% valign=middle>{logo}</td>
<td align=center valign=middle style="font-size:10pt">{head}</td>
<td width=15%></td></tr></table>
<h3 align=center style="margin-top:8px">DERS GÖREVLENDİRME ÇİZELGESİ</h3>
<p align=center style="font-size:10pt"><b>{e(date_tr(plan.date))}</b></p>
<table width=100% cellspacing=0 cellpadding=3 style="font-size:8.5pt">
<tr><td width=22%><b>Nöbetçi öğretmenler:</b></td><td>{duty}</td></tr></table>
<br>
{'<p style="margin:0">&nbsp;</p>'.join(tables) or '<p align=center>Görevlendirme yok</p>'}
<p style="font-size:8pt; color:#475569">Yukarıda adı geçen öğretmenler belirtilen ders saatlerinde
ilgili sınıflarda görevlendirilmiştir. Mor satırlar joker öğretmenlere, sarı satırlar nöbetçi
olmayan diğer öğretmenlere verilen derslerdir.</p>
<br><br>
<table width=100%><tr><td width=20%></td>{''.join(sign)}</tr></table>
</body></html>"""


def save_day_xlsx(path, plan, school=""):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Görevlendirme"
    ws.append([school])
    ws.append([f"Ders Görevlendirme Çizelgesi — {date_tr(plan.date)}"])
    ws.append([])
    ws.append(["Gelmeyen Öğretmen", "Mazeret", "Saat", "Sınıf", "Ders", "Derse Girecek Öğretmen",
               "Tür", "Not", "İmza"])
    for c in ws[4]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDE4EE")
    fills = {"bos": PatternFill("solid", fgColor="FFF4CC"), "yok": PatternFill("solid", fgColor="FFD6D6"),
             "joker": PatternFill("solid", fgColor="EFE2FF")}
    for a in plan.assignments:
        sinif, ders = split_info(a.info)
        ders = lesson_label(ders)
        ws.append([a.absent, plan.absences.get(a.absent, ""), f"{a.hour}. ders", sinif, ders or a.info,
                   a.assigned or "—", KIND_LABELS.get(a.kind, a.kind), a.note, ""])
        if a.kind in fills:
            for c in ws[ws.max_row]:
                c.fill = fills[a.kind]
    ws.append([])
    ws.append(["Gelmeyenler", ", ".join(f"{t} ({r})" for t, r in sorted(plan.absences.items()))])
    ws.append(["Nöbetçiler", ", ".join(t + (f" ({p})" if p else "") for t, p in plan.duty)])
    for col, w in zip("ABCDEFGHI", (24, 14, 9, 9, 18, 26, 22, 50, 16)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=5):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"].font = Font(bold=True)
    wb.save(path)


REPORT_HEADERS = ["Öğretmen", "Nöbet günü", "Bu ay", "Yıl toplamı", "Yıl: nöbetçiyken",
                  "Yıl: diğer", "Durum"]


def report_rows(month_counts, year_counts, duty_days):
    """Rapor satırları. Durum, nöbetçilerin yıllık ortalamasına göre yazılır."""
    teachers = sorted(set(month_counts) | set(year_counts) | set(duty_days), key=normalize)
    duty_year = [year_counts.get(t, (0, 0, 0))[0] for t in teachers if duty_days.get(t)]
    avg = sum(duty_year) / len(duty_year) if duty_year else 0
    rows = []
    for t in teachers:
        m_tot = month_counts.get(t, (0, 0, 0))[0]
        y_tot, y_duty, y_free = year_counts.get(t, (0, 0, 0))
        days = duty_days.get(t, 0)
        if not days:
            status = "Nöbetçi değil"
        elif y_tot < avg - 0.5:
            status = "Az girdi — öncelikli"
        elif y_tot > avg + 0.5:
            status = "Çok girdi — geride kalır"
        else:
            status = "Ortalamada"
        rows.append([t, days, m_tot, y_tot, y_duty, y_free, status])
    rows.sort(key=lambda r: (r[1] == 0, -r[3], normalize(r[0])))
    return rows


def save_report_xlsx(path, month, rows, school=""):
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Aylık Rapor"
    ws.append([school])
    ws.append([f"{month_tr(month)} Ders Doldurma Dağılımı (eğitim yılı toplamıyla)"])
    ws.append([])
    ws.append(REPORT_HEADERS)
    for c in ws[4]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDE4EE")
    for r in rows:
        ws.append(r)
    for col, w in zip("ABCDEFG", (28, 16, 18, 20, 18, 20, 38)):
        ws.column_dimensions[col].width = w
    wb.save(path)


RECORDS_FILE = "Kayıtlar.xlsx"


def write_records_xlsx(path, db, school=""):
    """Tüm kayıtları tek bir Excel dosyasına yazar (her kayıttan sonra yenilenir).

    Dosya Excel'de açıkken yazılamaz; o durumda PermissionError yükselir ve veriler
    programın veritabanında güvende kalır, bir sonraki kayıtta yeniden denenir.
    """
    import os
    import tempfile

    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    head_fill = PatternFill("solid", fgColor="DDE4EE")
    fills = {"bos": PatternFill("solid", fgColor="FFF4CC"),
             "yok": PatternFill("solid", fgColor="FFD6D6"),
             "joker": PatternFill("solid", fgColor="EFE2FF")}

    def sheet(wb, title, headers, widths, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = head_fill
        for i, w in enumerate(widths):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w
        ws.freeze_panes = "A2"
        return ws

    wb = openpyxl.Workbook()
    ws = sheet(wb, "Görevlendirmeler",
               ["Tarih", "Gün", "Saat", "Gelmeyen Öğretmen", "Mazeret", "Sınıf", "Ders",
                "Derse Giren Öğretmen", "Tür", "Not"], [12, 11, 6, 26, 12, 9, 18, 26, 24, 60],
               first=True)
    rows = db.conn.execute(
        "SELECT a.date, a.hour, a.absent_teacher, ab.reason, a.info, a.assigned, a.kind, a.note "
        "FROM assignments a LEFT JOIN absences ab ON ab.date = a.date "
        "AND ab.teacher = a.absent_teacher ORDER BY a.date, a.hour, a.absent_teacher").fetchall()
    for date, hour, absent, reason, info, assigned, kind, note in rows:
        d = dt.date.fromisoformat(date)
        sinif, ders = split_info(info)
        ders = lesson_label(ders)
        ws.append([d, DAYS[d.weekday()] if d.weekday() < 5 else "", hour, absent, reason or "",
                   sinif, ders or info or "", assigned or "—", KIND_LABELS.get(kind, kind or ""),
                   note or ""])
        ws.cell(ws.max_row, 1).number_format = "DD.MM.YYYY"
        if kind in fills:
            for c in ws[ws.max_row]:
                c.fill = fills[kind]
    for row in ws.iter_rows(min_row=2):
        row[-1].alignment = Alignment(wrap_text=True, vertical="top")

    ws = sheet(wb, "Devamsızlık", ["Tarih", "Gün", "Öğretmen", "Durumu"], [12, 11, 28, 12])
    for date, teacher, reason in db.conn.execute(
            "SELECT date, teacher, reason FROM absences ORDER BY date, teacher"):
        d = dt.date.fromisoformat(date)
        ws.append([d, DAYS[d.weekday()] if d.weekday() < 5 else "", teacher, reason or ""])
        ws.cell(ws.max_row, 1).number_format = "DD.MM.YYYY"

    ws = sheet(wb, "Aylık Özet", ["Ay"] + REPORT_HEADERS, [14, 28, 16, 20, 16, 18, 18])
    months = [r[0] for r in db.conn.execute(
        "SELECT DISTINCT substr(date, 1, 7) FROM assignments ORDER BY 1 DESC")]
    for m in months:
        last = max(db.saved_dates(m) or [m + "-28"])
        for r in report_rows(db.counts(m), db.year_counts(last), db.duty_days_in_month(m)):
            ws.append([month_tr(m)] + r)

    ws = sheet(wb, "Ders Programı", ["Geçerlilik", "Öğretmen", "Gün", "Saat", "Sınıf / Ders"],
               [12, 28, 11, 6, 30])
    for vid, valid_from, *_ in db.schedule_versions():
        sched = db.schedule(vid)
        for t in sorted(sched, key=normalize):
            for (d, h), info in sorted(sched[t].items()):
                ws.append([dt.date.fromisoformat(valid_from), t, DAYS[d], h, info])
                ws.cell(ws.max_row, 1).number_format = "DD.MM.YYYY"

    ws = sheet(wb, "Nöbet Çizelgesi", ["Ay", "Tarih", "Gün", "Öğretmen", "Nöbet Yeri"],
               [14, 12, 11, 28, 18])
    for m, _n in db.duty_months():
        for wd, date, t, p in sorted(db.duty_rows(m), key=lambda r: (r[1] or "", r[0], r[3] or "")):
            ws.append([month_tr(m), dt.date.fromisoformat(date) if date else "Her hafta",
                       DAYS[wd], t, p or ""])
            if date:
                ws.cell(ws.max_row, 2).number_format = "DD.MM.YYYY"

    if school:
        wb.properties.title = f"{school} — NöbetGörev kayıtları"
    folder = os.path.dirname(path) or "."
    os.makedirs(folder, exist_ok=True)
    fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=folder)
    os.close(fd)
    try:
        wb.save(tmp)
        os.replace(tmp, path)  # Excel'de açıksa burada PermissionError
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)


MARK = "\u2063"   # görünmez işaret: sayfalamada bölüm sınırlarını bulmak için


def duty_html(month, entries, info=None, has_logo=False, exempt=None, places=None, breaks=()):
    """Nöbetmatik çizelgesinin resmî çıktısı: her hafta için Günler × Nöbet yerleri tablosu.

    breaks: önüne sayfa sonu konacak bölümlerin sırası (0.. haftalar, son sıra imzalar).
    Her haftanın başlığı MARK ile başlar; imza bölümünün önünde ve arkasında da MARK vardır.
    """
    import datetime as _dt
    e = html.escape
    info = info or {}
    months = ["Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran", "Temmuz", "Ağustos",
              "Eylül", "Ekim", "Kasım", "Aralık"]
    weeks = {}
    for x in entries:
        if not x.date:
            continue
        d = _dt.date.fromisoformat(x.date)
        monday = d - _dt.timedelta(days=d.weekday())
        weeks.setdefault(monday, {}).setdefault(d, []).append(x)
    order = list(places or [])
    for x in entries:
        if x.place and x.place not in order and x.place != "N.İDARECİ":
            order.append(x.place)
    has_admin = any(x.place == "N.İDARECİ" for x in entries)
    cols = order + (["Nöbetçi İdareci"] if has_admin else [])
    blocks = []
    for monday in sorted(weeks):
        days = weeks[monday]
        first, last = min(days), max(days)
        title = (f"{first.day:02d} {months[first.month - 1].upper()}" if first == last else
                 f"{first.day:02d}-{last.day:02d} {months[last.month - 1].upper()}"
                 if first.month == last.month else
                 f"{first.day:02d} {months[first.month - 1].upper()} - "
                 f"{last.day:02d} {months[last.month - 1].upper()}")
        rows = []
        for d in sorted(days):
            cells = []
            for p in order:
                names = [x.teacher for x in days[d] if x.place == p]
                cells.append("<br>".join(e(n) for n in names) or "&nbsp;")
            if has_admin:
                cells.append("<br>".join(e(x.teacher) for x in days[d] if x.place == "N.İDARECİ")
                             or "&nbsp;")
            rows.append(f"<tr><td><b>{DAYS[d.weekday()].upper()}</b><br>"
                        f"<span style='color:#475569'>{d.strftime('%d.%m.%Y')}</span></td>"
                        + "".join(f"<td>{c}</td>" for c in cells) + "</tr>")
        # Sütun genişlikleri her haftada aynı olsun (tablolar alt alta düzgün hizalanır)
        place_w = (86 - (14 if has_admin else 0)) // max(1, len(order))
        head = "".join(f"<th width={14 if c == 'Nöbetçi İdareci' else place_w}%>{e(c.upper())}</th>"
                       for c in cols)
        brk = "; page-break-before:always" if len(blocks) in breaks else ""
        blocks.append(
            f"<p style='margin:6px 0 2px 0; font-size:8.5pt{brk}'>{MARK}<b>{title}</b></p>"
            "<table border=1 cellspacing=0 cellpadding=2 width=100% "
            "style='border-collapse:collapse; font-size:7.5pt'>"
            f"<tr style='background:#e2e8f0'><th width=14%>GÜNLER</th>{head}</tr>"
            + "".join(rows) + "</table>")
    sign_break = len(blocks) in breaks
    logo = "<img src='logo' height=64>" if has_logo else ""
    school = info.get("school", "")
    city = info.get("school_city", "")
    head = "<br>".join(x for x in ("T.C.", e(city), f"<b>{e(school)}</b>") if x)
    y, m = month.split("-")
    all_days = [_dt.date.fromisoformat(x.date) for x in entries if x.date]
    span = (f"{min(all_days).strftime('%d.%m.%Y')} – {max(all_days).strftime('%d.%m.%Y')}"
            if all_days else "")
    ex = ""
    if exempt:
        ex = ("<p style='font-size:8.5pt'><b>Nöbetten muaf öğretmenler:</b> " +
              ", ".join(f"{e(t)} ({e(r)})" if r else e(t) for t, r in sorted(exempt.items())) +
              "</p>")
    sign = []
    for title, key in (("Müdür Yardımcısı", "vice_principal"), ("Okul Müdürü", "principal")):
        name = (info.get(key, "") or "").replace(";", ",").split(",")[0].strip()
        sign.append(f"<td align=center width=40%>{e(name) or '&nbsp;'}<br><b>{title}</b>"
                    "<br><br>İmza</td>")
    return f"""
<html><body style="font-family:'Segoe UI', Arial; font-size:9pt; color:#111">
<table width=100% cellspacing=0 cellpadding=0><tr>
<td width=15% valign=middle>{logo}</td>
<td align=center valign=middle style="font-size:10pt">{head}</td>
<td width=15%></td></tr></table>
<h3 align=center style="margin:6px 0 0 0">{y} YILI {months[int(m) - 1].upper()} AYI NÖBET ÇİZELGESİ</h3>
<p align=center style="margin:2px 0 6px 0">{span}</p>
{''.join(blocks)}
{ex}
<p style="margin:0{'; page-break-before:always' if sign_break else ''}">{MARK}</p>
<table width=100%><tr><td width=20%></td>{''.join(sign)}</tr></table>
<p style="margin:0">{MARK}</p>
</body></html>"""


def fairness_summary(stats):
    """Adalet göstergesi: (karar, renk, açıklama satırları).

    Karar, kaydedilen her görevin kurala uygun olup olmadığına göre verilir: o saatte boş olan
    nöbetçiler arasından o ana kadar en az derse girmiş olan mı seçildi?
    """
    meta = stats.get("__meta__", {})
    people = {t: v for t, v in stats.items() if t != "__meta__"}
    duty = {t: v for t, v in people.items() if v["days"]}
    total = sum(v["assigned"] for v in people.values())
    if not total:
        return "Henüz görevlendirme yok", "#64748b", []
    exc = meta.get("exceptions", [])
    lines = []
    counts = [v["assigned"] for v in duty.values() if v["chances"]]
    if counts:
        lines.append(f"Nöbetçiler arasında en çok giren ile en az giren arasındaki fark: "
                     f"<b>{max(counts) - min(counts)} ders</b>.")
    load = meta.get("weekday_load", {})
    if load and len(counts) > 1 and max(counts) - min(counts) > 1:
        per = {d: n / max(1, k) for d, (n, k) in load.items()}
        hi = max(per, key=per.get)
        lo = min(per, key=per.get)
        if per[hi] > per[lo] * 1.3:
            lines.append(f"Fark gün yükünden: <b>{DAYS[hi]}</b> günleri nöbetçi başına "
                         f"{per[hi]:.1f} boş ders düştü, <b>{DAYS[lo]}</b> günleri {per[lo]:.1f}. "
                         "Bu gün yükü farkı, otomatik nöbet çizelgesinin sonraki aylardaki gün değişimiyle dengelenir.")
    if not exc:
        return ("✅ Adil — her görevde sırası gelen öğretmen seçildi", "#059669",
                [f"Kaydedilen <b>{total}</b> görevin hepsinde, o saatte boş olan nöbetçiler "
                 "arasından o ana kadar en az derse girmiş olan görevlendirildi."] + lines)
    names = ", ".join(sorted({x[2] for x in exc}, key=normalize))
    return (f"⚠ {len(exc)} görevde sıra atlandı (elle değişiklik)", "#d97706",
            [f"{total} görevin {total - len(exc)} tanesinde kural uygulandı; {len(exc)} görevde "
             f"elle seçim yapıldı: {names}. Program sonraki görevlendirmelerde az girenlere "
             "öncelik vererek farkı kapatır."] + lines)


def report_html(period_title, stats, info=None, has_logo=False):
    e = html.escape
    info = info or {}
    verdict, color, lines = fairness_summary(stats)
    rows = sorted(((t, v) for t, v in stats.items()
                   if t != "__meta__" and (v["days"] or v["assigned"])),
                  key=lambda x: (-x[1]["assigned"], normalize(x[0])))
    body = "".join(
        f"<tr><td>{i}</td><td>{e(t)}</td><td align=center>{v['days']}</td>"
        f"<td align=center>{v['chances']}</td><td align=center>{v['as_duty']}</td>"
        f"<td align=center>{v['as_other']}</td><td align=center><b>{v['assigned']}</b></td></tr>"
        for i, (t, v) in enumerate(rows, start=1))
    logo = "<img src='logo' height=60>" if has_logo else ""
    head = "<br>".join(x for x in ("T.C.", e(info.get("school_city", "")),
                                   f"<b>{e(info.get('school', ''))}</b>") if x)
    sign = []
    for title, key in (("Müdür Yardımcısı", "vice_principal"), ("Okul Müdürü", "principal")):
        name = (info.get(key, "") or "").replace(";", ",").split(",")[0].strip()
        sign.append(f"<td align=center width=40%>{e(name) or '&nbsp;'}<br><b>{title}</b>"
                    "<br><br>İmza</td>")
    return f"""
<html><body style="font-family:'Segoe UI', Arial; font-size:9pt; color:#111">
<table width=100%><tr><td width=15%>{logo}</td><td align=center style="font-size:10pt">{head}</td>
<td width=15%></td></tr></table>
<h3 align=center style="margin:6px 0 2px 0">BOŞ DERS GÖREVLENDİRME RAPORU</h3>
<p align=center style="margin:0 0 8px 0">{e(period_title)}</p>
<p style="font-size:11pt; color:{color}"><b>{e(verdict)}</b></p>
<p>{'<br>'.join(lines)}</p>
<p style="color:#475569">Görevlendirme kuralı: Bir öğretmen gelmediğinde her dersi için, o saatte
<b>kendi dersi olmayan</b> nöbetçi öğretmenler arasından <b>bu eğitim yılında en az derse girmiş
olan</b> seçilir. Boş nöbetçi yoksa o saatte dersi olmayan diğer öğretmenler arasından seçilir.
İdareciler otomatik görevlendirilmez.</p>
<table border=1 cellspacing=0 cellpadding=4 width=100% style="border-collapse:collapse">
<tr style="background:#e2e8f0"><th width=5%>No</th><th>Öğretmen</th><th>Nöbet tuttuğu gün</th>
<th>Derse girebileceği boş saat</th><th>Nöbetçiyken girdiği</th><th>Nöbetçi değilken girdiği</th>
<th>Toplam</th></tr>
{body}</table>
<p style="font-size:8pt; color:#475569">"Derse girebileceği boş saat": bir ders boşa düştüğünde
öğretmenin o gün nöbetçi olup o saatte kendi dersinin olmadığı durum sayısı.</p>
<br><table width=100%><tr><td width=20%></td>{''.join(sign)}</tr></table>
</body></html>"""
