"""Ana pencere ve diyaloglar (PySide6)."""
import datetime as dt
import os
import re
import shutil

from PySide6.QtCore import (QDate, QLocale, QObject, QRect, QSize, QStandardPaths, Qt, QThread,
                            QTimer, QUrl, Signal)
from PySide6.QtGui import QColor, QDesktopServices, QImage, QPageLayout, QPixmap, QTextDocument
from PySide6.QtWidgets import (QAbstractItemView, QApplication, QBoxLayout, QCheckBox, QComboBox,
                               QStackedWidget,
                               QFrame,
                               QCompleter, QDateEdit, QDialog, QDialogButtonBox, QFileDialog,
                               QFormLayout, QGroupBox, QHBoxLayout, QHeaderView, QInputDialog,
                               QLabel,
                               QLineEdit, QListWidget, QListWidgetItem, QMainWindow, QMessageBox, QPlainTextEdit,
                               QProgressDialog, QPushButton, QSpinBox, QSplitter, QStyledItemDelegate, QTableWidget,
                               QTableWidgetItem, QVBoxLayout, QWidget)

from . import export, theme, updater
from .theme import BLUE, GREEN, PURPLE, RED, YELLOW, page_header, variant
from .db import Database, data_dir
from .names import DAYS, NameMatcher, clean_name, lesson_label, normalize, split_info
from .parse_duty import DutyEntry, is_admin_place, parse_duty
from .parse_schedule import parse_schedule
from .reader import ReadError, read_document
from .scheduler import KIND_LABELS, options_for, plan_day, teachers_with_lessons
from .version import APP_NAME, APP_TAGLINE, __version__

STATUSES = ["Okulda", "Raporlu", "İzinli", "Görevli", "Sevkli", "Kurs / Seminer", "Diğer…"]
OTHER = "Diğer…"
FILE_FILTER = "Belgeler (*.pdf *.xlsx *.xlsm *.csv);;PDF (*.pdf);;Excel (*.xlsx *.xlsm)"


def _item(text, editable=False):
    it = QTableWidgetItem("" if text is None else str(text))
    if not editable:
        it.setFlags(it.flags() & ~Qt.ItemIsEditable)
    return it


def fill_duty_grid(table, headers, rows):
    """Nöbet çizelgesi tablosunu düzenli çizer.

    rows: [(gün etiketi, [[isim, ...], [isim, ...], ...]), ...] — her hücrede isimler alt alta.
    Gün sütunu sabit genişlikte, yer sütunları eşit genişlikte; isimler sola-yukarı hizalı.
    """
    from PySide6.QtGui import QFontMetrics
    table.clear()
    table.setColumnCount(len(headers))
    table.setRowCount(len(rows))
    table.setHorizontalHeaderLabels(headers)
    table.setShowGrid(True)
    table.setAlternatingRowColors(False)
    table.setWordWrap(True)
    table.setTextElideMode(Qt.ElideNone)
    table.verticalHeader().hide()
    table.setSelectionMode(QAbstractItemView.NoSelection)
    table.setFocusPolicy(Qt.NoFocus)
    h = table.horizontalHeader()
    h.setSectionResizeMode(0, QHeaderView.Fixed)
    table.setColumnWidth(0, 118)
    for c in range(1, len(headers)):
        h.setSectionResizeMode(c, QHeaderView.Stretch)
    if headers and headers[-1] == "Nöbetçi idareci":
        h.setSectionResizeMode(len(headers) - 1, QHeaderView.Fixed)
        table.setColumnWidth(len(headers) - 1, 150)
    h.setDefaultAlignment(Qt.AlignLeft | Qt.AlignVCenter)
    line = QFontMetrics(table.font()).lineSpacing()
    for r, (label, cells) in enumerate(rows):
        day = _item(label)
        f = day.font()
        f.setBold(True)
        day.setFont(f)
        day.setTextAlignment(Qt.AlignLeft | Qt.AlignTop)
        day.setBackground(BLUE)
        table.setItem(r, 0, day)
        most = max([2] + [len(names) for names in cells])
        for c, names in enumerate(cells, start=1):
            it = _item("\n".join(names))
            it.setTextAlignment(Qt.AlignLeft | Qt.AlignTop)
            table.setItem(r, c, it)
        table.setRowHeight(r, most * line + 18)
    # Satır yüksekliğini sütunlar son genişliğini aldıktan sonra içeriğe göre ayarla; pencere
    # boyutu değişince yeniden hesapla (uzun isim kesilmesin, satırlar da gereksiz uzamasın).
    def fit():
        table.resizeRowsToContents()
        for r in range(table.rowCount()):
            table.setRowHeight(r, table.rowHeight(r) + 6)
    fit()
    QTimer.singleShot(0, fit)
    if not getattr(table, "_fit_hooked", False):
        table._fit_hooked = True
        h.sectionResized.connect(lambda *_: QTimer.singleShot(0, table._fit))
    table._fit = fit


def _qdate(iso):
    return QDate.fromString(iso, "yyyy-MM-dd")


def _iso(qdate):
    return qdate.toString("yyyy-MM-dd")


def _message(parent, icon, title, text, buttons):
    """Türkçe düğmeli ileti kutusu; basılan düğmenin sırasını döner."""
    box = QMessageBox(icon, title, text, QMessageBox.NoButton, parent)
    made = []
    for label, role in buttons:
        made.append(box.addButton(label, role))
    box.setDefaultButton(made[0])
    box.exec()
    clicked = box.clickedButton()
    box.deleteLater()
    return made.index(clicked) if clicked in made else len(made) - 1


def _info(parent, text, title="Bilgi"):
    _message(parent, QMessageBox.Information, title, text, [("Tamam", QMessageBox.AcceptRole)])


def _warn(parent, title, text):
    _message(parent, QMessageBox.Warning, title, text, [("Tamam", QMessageBox.AcceptRole)])


def _ask(parent, text, title="Onay"):
    return _message(parent, QMessageBox.Question, title, text,
                    [("Evet", QMessageBox.YesRole), ("Hayır", QMessageBox.NoRole)]) == 0


class ChoiceButton(QPushButton):
    """Tablo satırları için hafif açılır seçim (QComboBox yerine).

    Her QComboBox arka planda ayrı bir açılır pencere taşır; uzun listelerde bu, tema
    değişimini ve sayfa açılışını yavaşlatır. Menü yalnızca tıklanınca oluşturulur.
    QComboBox'ın burada kullanılan kısmıyla aynı arayüze sahiptir.
    """

    currentTextChanged = Signal(str)

    def __init__(self, items=(), data=None):
        super().__init__()
        self.items = list(items)
        self.data = list(data) if data is not None else None
        self._text = self.items[0] if self.items else ""
        self.setCursor(Qt.PointingHandCursor)
        self.setObjectName("choice")
        self.clicked.connect(self._open)
        self._show()

    def _show(self):
        self.setText(f"{self._text}   ▾")

    def _open(self):
        from PySide6.QtWidgets import QMenu
        menu = QMenu(self)
        for it in self.items:
            act = menu.addAction(it)
            act.setCheckable(True)
            act.setChecked(it == self._text)
        chosen = menu.exec(self.mapToGlobal(self.rect().bottomLeft()))
        menu.deleteLater()
        if chosen is not None:
            self.setCurrentText(chosen.text())

    def addItems(self, items):
        self.items += list(items)

    def currentText(self):
        return self._text

    def currentData(self):
        if self.data is None or self._text not in self.items:
            return None
        return self.data[self.items.index(self._text)]

    def setCurrentText(self, text):
        if text == self._text:
            return
        self._text = text
        self._show()
        if not self.signalsBlocked():
            self.currentTextChanged.emit(text)

    def findText(self, text):
        return self.items.index(text) if text in self.items else -1

    def count(self):
        return len(self.items)

    def insertItem(self, index, text):
        self.items.insert(index, text)


class _AlignDelegate(QStyledItemDelegate):
    """Sütun hizası: metin sütunları sola, sayı/tarih sütunları ortaya; hepsi dikeyde ortalı."""

    def __init__(self, parent, center):
        super().__init__(parent)
        self.center = set(center)

    def initStyleOption(self, option, index):
        super().initStyleOption(option, index)
        table = self.parent()
        spanned = table.columnSpan(index.row(), index.column()) > 1  # öğretmen başlık satırı
        option.displayAlignment = (Qt.AlignCenter if index.column() in self.center and not spanned
                                   else Qt.AlignLeft | Qt.AlignVCenter)


def align(table, *center):
    """Tablonun başlıklarını ve hücrelerini aynı hizaya getirir (başlık, altındaki yazıyla
    aynı yerden başlar). center: ortalanacak sütunlar."""
    table.setItemDelegate(_AlignDelegate(table, center))
    header = table.horizontalHeader()
    header.setDefaultAlignment(Qt.AlignLeft | Qt.AlignVCenter)
    for c in range(table.columnCount()):
        it = table.horizontalHeaderItem(c)
        if it is not None:
            it.setTextAlignment(Qt.AlignCenter if c in center else Qt.AlignLeft | Qt.AlignVCenter)
    return table


class HourStripDelegate(QStyledItemDelegate):
    """Bir günün ders saatlerini numaralı küçük kutucuklarla çizer (pencere açmadan, hızlı).

    Hücre verisi (Qt.UserRole): {"max": 8, "lessons": {1, 2}} — dersi olan saatler dolu;
    ya da {"max": 8, "need": {3, 5}, "busy": {5}} — boş kalan saatlerde yeşil = girebilir,
    kırmızı = kendi dersi var; diğer saatler soluk.
    """
    BOX, GAP, LABEL = 24, 5, 14

    def paint(self, painter, option, index):
        from PySide6.QtGui import QPen
        from PySide6.QtWidgets import QStyle, QStyleOptionViewItem
        opt = QStyleOptionViewItem(option)
        self.initStyleOption(opt, index)
        opt.text = ""
        opt.widget.style().drawControl(QStyle.CE_ItemViewItem, opt, painter, opt.widget)
        data = index.data(Qt.UserRole) or {}
        hours = (list(range(1, int(data.get("max", 8)) + 1)) if "lessons" in data
                 else sorted(data.get("need", ())))
        dark = bool(theme._dark)
        painter.save()
        painter.setRenderHint(painter.RenderHint.Antialiasing)
        f = painter.font()
        f.setPointSizeF(max(7.0, f.pointSizeF() - 1.5))
        f.setBold(True)
        painter.setFont(f)
        labeled = True                    # kutunun altına "boş" / "dolu" yazılır
        x = option.rect.left() + 8
        total = self.BOX + (self.LABEL if labeled else 0)
        y = option.rect.top() + (option.rect.height() - total) // 2
        small = painter.font()
        small.setBold(False)
        small.setPointSizeF(max(6.5, small.pointSizeF() - 1))
        bold = painter.font()
        muted = QColor("#64748b" if not dark else "#94a3b8")
        for h in hours:
            r = QRect(x, y, self.BOX, self.BOX)
            if "lessons" in data:
                on = h in data["lessons"]
                fill = QColor(theme.ACCENT) if on else QColor(0, 0, 0, 0)
                pen = QColor(theme.ACCENT) if on else QColor("#cbd5e1" if not dark else "#475569")
                text = QColor("white") if on else muted
            else:
                need, busy = data.get("need", set()), data.get("busy", set())
                if h in need and h not in busy:
                    fill, pen, text = QColor("#16a34a"), QColor("#16a34a"), QColor("white")
                elif h in need:
                    fill, pen, text = (QColor("#fee2e2" if not dark else "#47201f"),
                                       QColor("#f87171"), QColor("#b91c1c" if not dark
                                                                 else "#fecaca"))
                else:
                    fill, pen = QColor(0, 0, 0, 0), QColor("#e2e8f0" if not dark else "#334155")
                    text = QColor("#cbd5e1" if not dark else "#475569")
            painter.setPen(QPen(pen, 1.2))
            painter.setBrush(fill)
            painter.drawRoundedRect(r, 6, 6)
            painter.setPen(text)
            painter.setFont(bold)
            painter.drawText(r, Qt.AlignCenter, str(h))
            if labeled:
                if "lessons" in data:
                    free = h not in data["lessons"]
                    color = muted if free else QColor(theme.ACCENT)
                else:
                    free = h not in data.get("busy", set())
                    color = QColor("#15803d") if free else QColor("#b91c1c" if not dark
                                                                  else "#fca5a5")
                painter.setFont(small)
                painter.setPen(color)
                painter.drawText(QRect(x - 4, y + self.BOX + 1, self.BOX + 8, self.LABEL),
                                 Qt.AlignHCenter | Qt.AlignTop, "boş" if free else "dolu")
            x += self.BOX + self.GAP
        painter.restore()

    def sizeHint(self, option, index):
        data = index.data(Qt.UserRole) or {}
        n = int(data.get("max", 8)) if "lessons" in data else len(data.get("need", ()))
        extra = self.LABEL
        return QSize(16 + n * (self.BOX + self.GAP), self.BOX + 14 + extra)


class _StillFilter(QObject):
    """Fare yalnızca gezdirildiğinde (tık / tekerlek olmadan) tabloda hiçbir şey değişmesin:
    ipucu kutuları açılmaz, üzerine gelme vurgusu olmaz."""

    def eventFilter(self, obj, event):
        from PySide6.QtCore import QEvent
        return event.type() in (QEvent.ToolTip, QEvent.HoverEnter, QEvent.HoverMove,
                                QEvent.HoverLeave)


_still = None


def keep_still(*views):
    global _still
    if _still is None:
        _still = _StillFilter()
    for v in views:
        if hasattr(v, "setAutoScroll"):
            v.setAutoScroll(False)          # fare kenara gelince liste kendiliğinden kaymasın
        if hasattr(v, "setDragEnabled"):
            v.setDragEnabled(False)
        target = v.viewport() if hasattr(v, "viewport") else v
        target.setMouseTracking(False)
        target.setAttribute(Qt.WA_Hover, False)
        target.installEventFilter(_still)
        v.installEventFilter(_still)


def fit_day_document(plan, info, has_logo, schedule):
    """Görevlendirme belgesini ana çizelge (imzalar dahil) mümkün olduğunca TEK sayfaya
    sığacak şekilde hazırlar: sığmazsa yazı ve boşluklar adım adım küçültülür. İmzalar hep
    tabloların hemen altında kalır; öğretmen programları sonraki sayfadan başlar."""
    chosen = None
    for scale in (1.0, 0.92, 0.85, 0.78, 0.7, 0.62, 0.55):
        doc = keep_together(lambda breaks, sc=scale: export.day_html(
            plan, info, has_logo, schedule, breaks, sc))
        layout = doc.documentLayout()
        doc.pageCount()
        page = doc.pageSize().height()
        end = None
        block = doc.begin()
        while block.isValid():
            if export.SECTION_END in block.text():
                end = layout.blockBoundingRect(block).bottom()
                break
            block = block.next()
        pages = int(((end or 0) - 1) // page) + 1 if end else 1
        if chosen is None or pages < chosen[0]:
            chosen = (pages, doc)
        if pages == 1:
            break
    return chosen[1]


def run_dialog(dlg):
    """Pencereyi açar; kapanınca belleği bırakır (açık kalan eski pencereler programı ve tema
    değişimini yavaşlatır). Silme olay döngüsüne dönülünce yapılır, sonuçlar okunabilir."""
    ok = dlg.exec() == QDialog.Accepted
    dlg.deleteLater()
    return ok


def _stretch(table, *cols):
    # Fare / dokunmatik fare kenara yaklaşınca tablo kendiliğinden kaymasın
    table.setAutoScroll(False)
    table.setAlternatingRowColors(True)
    table.verticalHeader().setDefaultSectionSize(theme.ROW_HEIGHT)
    table.setShowGrid(False)
    h = table.horizontalHeader()
    h.setSectionResizeMode(QHeaderView.ResizeToContents)
    for c in cols:
        h.setSectionResizeMode(c, QHeaderView.Stretch)


class Worker(QThread):
    done = Signal(object)
    failed = Signal(str)
    progress = Signal(int, int)

    def __init__(self, fn, *args, with_progress=False):
        super().__init__()
        self.fn, self.args, self.with_progress = fn, args, with_progress

    def run(self):
        try:
            if self.with_progress:
                res = self.fn(*self.args, progress=lambda d, t: self.progress.emit(d, t))
            else:
                res = self.fn(*self.args)
            self.done.emit(res)
        except Exception as e:  # noqa: BLE001 - kullanıcıya gösterilecek
            self.failed.emit(str(e))


def needs_class(info):
    """Sınıfı olması gereken ama yazılmamış ders mi? "DESTEK" dersleri belgede sınıfsız
    yazılır; olduğu gibi bırakılır, eksik sayılmaz."""
    return bool(info) and not split_info(info)[0] and "destek" not in normalize(info)


def missing_class(data):
    """Sınıfı okunamayan dersler: {öğretmen: [(gün, saat, bilgi), ...]}."""
    out = {}
    for t, slots in data.items():
        for (d, h), info in sorted(slots.items()):
            if needs_class(info):
                out.setdefault(t, []).append((d, h, info))
    return out


def missing_class_text(data):
    miss = missing_class(data)
    if not miss:
        return ""
    n = sum(len(v) for v in miss.values())
    kinds = sorted({i for v in miss.values() for _d, _h, i in v})
    return (f"⚠ <b>{n} derste sınıf yazmıyor</b> ({len(miss)} öğretmen; ör. "
            f"{', '.join(kinds[:3])}). PDF'te bu hücrelerde sınıf adı yok. Sarı hücrelere çift "
            "tıklayıp sınıfı ekleyebilirsiniz (ör. <i>6/A BİLG.</i>).")


class WeekGrid(QTableWidget):
    """Bir öğretmenin haftalık programı (satır: saat, sütun: gün)."""

    def __init__(self, hours=8, editable=False):
        super().__init__(hours, 5)
        self.editable = editable
        self.setHorizontalHeaderLabels(DAYS)
        self.setVerticalHeaderLabels([f"{h}. ders" for h in range(1, hours + 1)])
        self.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.setWordWrap(True)
        if editable:
            self.itemChanged.connect(self._paint)

    def _paint(self, it):
        """Hücre rengi içeriğe göre: ders + sınıf yeşil, sınıfı eksik sarı, boş renksiz."""
        text = (it.text() or "").strip()
        self.blockSignals(True)
        if not text:
            it.setBackground(QColor(0, 0, 0, 0))
            it.setToolTip("")
        elif needs_class(text):
            it.setBackground(YELLOW)
            it.setToolTip("Sınıf yazmıyor — çift tıklayıp sınıfı ekleyin (ör. 6/A BİLG.)")
        else:
            it.setBackground(GREEN)
            it.setToolTip("")
        self.blockSignals(False)

    def load(self, slots, hours=None):
        hours = max(hours or 8, max((h for _, h in slots), default=0))
        self.clearContents()
        self.setRowCount(hours)
        self.setVerticalHeaderLabels([f"{h}. ders" for h in range(1, hours + 1)])
        for (d, h), info in slots.items():
            if 0 <= d < 5 and h >= 1:
                it = _item(info, self.editable)
                if needs_class(info):
                    it.setBackground(YELLOW)   # sınıf okunamadı
                    it.setToolTip("Sınıf yazmıyor — çift tıklayıp sınıfı ekleyin (ör. 6/A BİLG.)")
                else:
                    it.setBackground(GREEN)
                self.setItem(h - 1, d, it)
        for r in range(self.rowCount()):
            for c in range(5):
                if not self.item(r, c):
                    self.setItem(r, c, _item("", self.editable))

    def slots(self):
        out = {}
        for r in range(self.rowCount()):
            for c in range(5):
                it = self.item(r, c)
                if it and it.text().strip():
                    out[(c, r + 1)] = it.text().strip()
        return out


# ---------------------------------------------------------------- diyaloglar

class ScheduleReviewDialog(QDialog):
    """Okunan ders programını kaydetmeden önce gösterir; isim düzeltme/birleştirme yapılabilir."""

    def __init__(self, parent, result, source, valid_from):
        super().__init__(parent)
        self.setWindowTitle("Ders programı önizleme — kontrol edip kaydedin")
        self.resize(1100, 700)
        self.result = result
        self.names = sorted(result.teachers, key=normalize)

        self.table = QTableWidget(len(self.names), 3)
        self.table.setHorizontalHeaderLabels(["Al", "Öğretmen (düzeltilebilir)", "Ders"])
        for r, n in enumerate(self.names):
            chk = _item("")
            chk.setFlags(Qt.ItemIsUserCheckable | Qt.ItemIsEnabled | Qt.ItemIsSelectable)
            chk.setCheckState(Qt.Checked)
            self.table.setItem(r, 0, chk)
            it = _item(n, editable=True)
            if n.startswith("Bilinmeyen"):
                it.setBackground(RED)
            self.table.setItem(r, 1, it)
            self.table.setItem(r, 2, _item(len(result.teachers[n])))
        _stretch(self.table, 1)
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.table.currentCellChanged.connect(self._show)

        self.grid = WeekGrid(result.max_hour, editable=True)
        self.grid.itemChanged.connect(self._grid_edited)
        self.current = None
        self.date = QDateEdit(_qdate(valid_from))
        self.date.setCalendarPopup(True)
        self.date.setDisplayFormat("dd.MM.yyyy")

        miss = missing_class(result.teachers)
        extra = []
        if miss:
            extra.append(f"{sum(len(v) for v in miss.values())} derste sınıf yazmıyor (sarı "
                         "hücreler): " + "; ".join(
                             f"{t}: " + ", ".join(f"{DAYS[d][:3]} {h}. {i}" for d, h, i in v)
                             for t, v in sorted(miss.items())))
        for r, n in enumerate(self.names):
            if n in miss:
                self.table.item(r, 1).setBackground(YELLOW)
                self.table.item(r, 1).setToolTip("Bazı derslerinde sınıf yazmıyor")
        warn = QPlainTextEdit("\n".join(list(result.warnings) + extra) or "Uyarı yok.")
        warn.setReadOnly(True)
        warn.setMaximumHeight(110)
        layouts = ", ".join(sorted(set(result.layouts))) or "—"

        left = QWidget()
        lv = QVBoxLayout(left)
        info_lbl = QLabel(f"<b>{len(self.names)} öğretmen bulundu.</b> Algılanan yerleşim: "
                            f"{layouts}<br>Yanlış okunan isimleri çift tıklayıp düzeltin. Aynı "
                            "ada çevrilen satırlar birleştirilir.<br><b>Sarı</b> öğretmen / "
                            "hücre: sınıf yazmıyor — sağdaki programda hücreye çift tıklayıp "
                            "sınıfı ekleyin (ör. <i>6/A BİLG.</i>); düzelince sarı kalkar. "
                            "<i>DESTEK</i> dersleri belgede sınıfsız yazılır; olduğu gibi kalır.")
        info_lbl.setWordWrap(True)
        lv.addWidget(info_lbl)
        lv.addWidget(self.table)
        split = QSplitter()
        split.addWidget(left)
        split.addWidget(self.grid)
        self.grid.setMinimumWidth(520)
        split.setSizes([380, 720])

        form = QFormLayout()
        form.addRow("Program şu tarihten itibaren geçerli:", self.date)
        form.addRow(QLabel(f"Kaynak: {os.path.basename(source)}"))
        buttons = QDialogButtonBox(QDialogButtonBox.Save | QDialogButtonBox.Cancel)
        buttons.button(QDialogButtonBox.Save).setText("Kaydet")
        buttons.button(QDialogButtonBox.Cancel).setText("Vazgeç")
        buttons.accepted.connect(self.accept)
        buttons.rejected.connect(self.reject)

        v = QVBoxLayout(self)
        v.addWidget(split, 1)
        v.addWidget(QLabel("Uyarılar:"))
        v.addWidget(warn)
        v.addLayout(form)
        v.addWidget(buttons)
        if self.names:
            self.table.selectRow(0)
            self._show(0, 0, -1, -1)

    def _show(self, row, *_):
        if 0 <= row < len(self.names):
            self.current = None
            self.grid.blockSignals(True)
            self.grid.load(self.result.teachers[self.names[row]], self.result.max_hour)
            self.grid.blockSignals(False)
            self.current = row

    def _grid_edited(self, _it):
        """Önizlemede düzeltilen dersi (ör. eksik sınıf) okunan programa yazar."""
        if self.current is None:
            return
        name = self.names[self.current]
        self.result.teachers[name] = self.grid.slots()
        miss = missing_class({name: self.result.teachers[name]})
        cell = self.table.item(self.current, 1)
        cell.setBackground(YELLOW if miss else QColor(0, 0, 0, 0))
        cell.setToolTip("Bazı derslerinde sınıf yazmıyor" if miss else "")
        self.table.item(self.current, 2).setText(str(len(self.result.teachers[name])))

    def teachers(self):
        out = {}
        for r, orig in enumerate(self.names):
            if self.table.item(r, 0).checkState() != Qt.Checked:
                continue
            name = clean_name(self.table.item(r, 1).text()) or orig
            slots = out.setdefault(name, {})
            for k, v in self.result.teachers[orig].items():
                slots[k] = f"{slots[k]} + {v}" if k in slots and slots[k] != v else v
        return out

    def valid_from(self):
        return _iso(self.date.date())


SCHOOL_KEYS = [("school", "Okul adı", "Ör. Atatürk Anadolu Lisesi"),
               ("school_city", "İl / İlçe", "Ör. BOLU / MERKEZ"),
               ("principal", "Okul müdürü", "Ad Soyad"),
               ("vice_principal", "Müdür yardımcıları", "Ad Soyad, Ad Soyad (virgülle ayırın)")]


def admin_roles_text(db):
    """'Okul Müdürü: X · Müdür Yardımcısı: Y, Z · Diğer idareci: W' biçiminde özet."""
    principal = db.get_setting("principal", "").strip()
    vices = [n.strip() for n in re.split(r"[,;\n]", db.get_setting("vice_principal", ""))
             if n.strip()]
    key = {normalize(n) for n in [principal] + vices if n}
    others = sorted((a for a in db.admins() if normalize(a) not in key), key=normalize)
    parts = []
    if principal:
        parts.append(f"Okul Müdürü: <b>{principal}</b>")
    if vices:
        parts.append(f"Müdür Yardımcısı: <b>{', '.join(vices)}</b>")
    if others:
        parts.append(f"Diğer idareci: <b>{', '.join(others)}</b>")
    return "  ·  ".join(parts)


PRINT_MARGIN_MM = 12


def _page_size_px(landscape=False):
    """A4 yazdırma alanı (kenar boşlukları düşülmüş), belge düzeninin birimiyle (piksel)."""
    from PySide6.QtCore import QSizeF
    w, h = 210 - 2 * PRINT_MARGIN_MM, 297 - 2 * PRINT_MARGIN_MM
    if landscape:
        w, h = h, w
    dpi = QApplication.primaryScreen().logicalDotsPerInchY() if QApplication.primaryScreen() \
        else 96
    return QSizeF(w * dpi / 25.4, h * dpi / 25.4)


def html_document(html, landscape=False):
    """Logolu, A4'e göre sayfalanmış yazdırma belgesi (önizleme, yazıcı ve PDF aynı görünür)."""
    doc = QTextDocument()
    if os.path.exists(logo_path()):
        doc.addResource(QTextDocument.ImageResource, QUrl("logo"), QImage(logo_path()))
    doc.setPageSize(_page_size_px(landscape))
    doc.setDocumentMargin(0)
    doc.setHtml(html)
    return doc


def keep_together(make_html, landscape=False):
    """Bölümleri (ör. haftalık nöbet tabloları) iki sayfaya bölmeden yerleştirir.

    make_html(breaks) HTML üretir; her bölüm export.MARK ile başlar, son MARK belgenin sonudur.
    Bir sayfaya sığmayan bölümden önce sayfa sonu konur ve belge yeniden ölçülür.
    """
    breaks = set()
    for _ in range(40):
        doc = html_document(make_html(breaks), landscape)
        layout = doc.documentLayout()
        doc.pageCount()  # düzeni tamamlat (yoksa konumlar eksik hesaplanır)
        page = doc.pageSize().height()
        # Üst düzeydeki bloklar ve tablolar sırayla: her MARK yeni bir bölüm başlatır
        spans = []           # [üst, alt] her bölüm için
        it = doc.rootFrame().begin()
        while not it.atEnd():
            frame, block = it.currentFrame(), it.currentBlock()
            if frame is not None:
                rect = layout.frameBoundingRect(frame)
            else:
                rect = layout.blockBoundingRect(block)
                if export.MARK in block.text():
                    spans.append([rect.top(), rect.bottom()])
            if spans and rect.height() > 0:
                spans[-1][1] = max(spans[-1][1], rect.bottom())
            it += 1
        moved = False
        for i, (top, bottom) in enumerate(spans[:-1]):
            if i not in breaks and int(top // page) != int((bottom - 1) // page) \
                    and bottom - top < page:
                breaks.add(i)
                moved = True
                break
        if not moved:
            return doc
    return doc


def preview(parent, doc, pdf_name, landscape=False):
    """Baskı önizleme: araç çubuğunda Yazdır'ın yanında 'PDF olarak kaydet' düğmesi de vardır."""
    from PySide6.QtGui import QAction, QPageSize
    from PySide6.QtPrintSupport import QPrinter, QPrintPreviewDialog
    from PySide6.QtWidgets import QToolBar
    orientation = QPageLayout.Landscape if landscape else QPageLayout.Portrait
    from PySide6.QtCore import QMarginsF
    margins = QMarginsF(*[PRINT_MARGIN_MM] * 4)
    printer = QPrinter(QPrinter.HighResolution)
    printer.setPageSize(QPageSize(QPageSize.A4))
    printer.setPageOrientation(orientation)
    printer.setPageMargins(margins, QPageLayout.Millimeter)
    dlg = QPrintPreviewDialog(printer, parent)
    dlg.setWindowTitle("Baskı önizleme")
    dlg.resize(1100, 800)
    dlg.paintRequested.connect(doc.print_)

    def save_pdf():
        docs = QStandardPaths.writableLocation(QStandardPaths.DocumentsLocation) or ""
        path, _ = QFileDialog.getSaveFileName(dlg, "PDF olarak kaydet",
                                              os.path.join(docs, pdf_name), "PDF (*.pdf)")
        if not path:
            return
        if not path.lower().endswith(".pdf"):
            path += ".pdf"
        pdf = QPrinter(QPrinter.HighResolution)
        pdf.setOutputFormat(QPrinter.PdfFormat)
        pdf.setOutputFileName(path)
        pdf.setPageSize(QPageSize(QPageSize.A4))
        pdf.setPageOrientation(orientation)
        pdf.setPageMargins(margins, QPageLayout.Millimeter)
        doc.print_(pdf)
        QDesktopServices.openUrl(QUrl.fromLocalFile(path))

    bar = dlg.findChild(QToolBar)
    if bar is not None:
        bar.addSeparator()
        act = QAction("📄  PDF olarak kaydet", dlg)
        act.triggered.connect(save_pdf)
        bar.addAction(act)
        btn = bar.widgetForAction(act)
        if btn is not None:
            btn.setToolButtonStyle(Qt.ToolButtonTextOnly)
            btn.setStyleSheet("font-weight:600; padding:4px 10px")
    dlg.exec()
    dlg.deleteLater()


def logo_path():
    return os.path.join(data_dir(), "logo.png")


def round_logo(img, size=512):
    """Logoyu kendi şekline göre hazırlar: kenarlardaki beyaz boşluk kırpılır.

    Logo yuvarlaksa (köşeleri boşsa) daire olarak kesilir ve köşelerdeki beyazlık saydam olur;
    kare veya dikdörtgen logolar kendi şekliyle kalır.
    """
    from PySide6.QtGui import QBrush, QColor, QPainter, QPainterPath
    img = img.convertToFormat(QImage.Format_ARGB32)
    # Kenarlardaki beyaz / saydam boşluğu kırp
    w, h = img.width(), img.height()

    def blank(x, y):
        c = img.pixelColor(x, y)
        return c.alpha() < 20 or (c.red() > 235 and c.green() > 235 and c.blue() > 235)

    step = max(1, min(w, h) // 200)
    xs = [x for x in range(0, w, step) if not all(blank(x, y) for y in range(0, h, step))]
    ys = [y for y in range(0, h, step) if not all(blank(x, y) for x in range(0, w, step))]
    if xs and ys:
        img = img.copy(min(xs), min(ys), max(xs) - min(xs) + 1, max(ys) - min(ys) + 1)
    w, h = img.width(), img.height()
    # Köşelere yakın noktalar boşsa logo yuvarlaktır
    inset = 0.06
    corners = [(int(w * inset), int(h * inset)), (int(w * (1 - inset)) - 1, int(h * inset)),
               (int(w * inset), int(h * (1 - inset)) - 1),
               (int(w * (1 - inset)) - 1, int(h * (1 - inset)) - 1)]
    is_round = 0.85 < w / max(1, h) < 1.18 and all(blank(x, y) for x, y in corners)
    if not is_round:
        return img.scaled(size, size, Qt.KeepAspectRatio, Qt.SmoothTransformation)
    side = min(img.width(), img.height())
    img = img.copy((img.width() - side) // 2, (img.height() - side) // 2, side, side)
    img = img.scaled(size, size, Qt.IgnoreAspectRatio, Qt.SmoothTransformation)
    out = QImage(size, size, QImage.Format_ARGB32)
    out.fill(Qt.transparent)
    p = QPainter(out)
    p.setRenderHint(QPainter.Antialiasing)
    path = QPainterPath()
    path.addEllipse(0, 0, size, size)
    p.setClipPath(path)
    p.drawImage(0, 0, img)
    p.end()
    return out


MONTHS_TR = ["Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran", "Temmuz", "Ağustos", "Eylül",
             "Ekim", "Kasım", "Aralık"]


class MonthPicker(QWidget):
    """Ay ve yıl açılır listeleri (yazmadan seçmek için)."""

    changed = Signal()

    def __init__(self, month=None):
        super().__init__()
        today = dt.date.today()
        self.m = QComboBox()
        self.m.addItems(MONTHS_TR)
        self.y = QComboBox()
        for yr in range(today.year - 2, today.year + 3):
            self.y.addItem(str(yr))
        for c in (self.m, self.y):
            c.setMinimumHeight(34)
            c.currentIndexChanged.connect(lambda _i: self.changed.emit())
        self.m.setMinimumWidth(140)
        h = QHBoxLayout(self)
        h.setContentsMargins(0, 0, 0, 0)
        h.addWidget(self.m)
        h.addWidget(self.y)
        h.addStretch()
        self.set(month or today.strftime("%Y-%m"))

    def set(self, month):
        y, m = month.split("-")
        if self.y.findText(y) < 0:
            self.y.addItem(y)
        self.m.blockSignals(True)
        self.y.blockSignals(True)
        self.m.setCurrentIndex(int(m) - 1)
        self.y.setCurrentText(y)
        self.m.blockSignals(False)
        self.y.blockSignals(False)

    def key(self):
        return f"{self.y.currentText()}-{self.m.currentIndex() + 1:02d}"


class SchoolTab(QWidget):
    """Okul adı, yöneticiler ve logo: çıktılarda ve sol menüde kullanılır."""

    def __init__(self, win):
        super().__init__()
        self.win, self.db = win, win.db
        form = QFormLayout()
        form.setVerticalSpacing(12)
        form.setHorizontalSpacing(16)
        self.fields = {}
        for key, label, hint in SCHOOL_KEYS[:2]:  # idareciler ders programından seçilir
            e = QLineEdit()
            e.setPlaceholderText(hint)
            e.setMinimumHeight(36)
            self.fields[key] = e
            form.addRow(label + ":", e)
        self.logo = QLabel()
        self.logo.setFixedSize(132, 132)
        self.logo.setAlignment(Qt.AlignCenter)
        self.logo.setObjectName("logoBox")
        self.new_logo = None
        self.remove_logo = False
        pick = QPushButton("🖼  Logo seç…")
        pick.clicked.connect(self._pick)
        clear = QPushButton("Logoyu kaldır")
        clear.clicked.connect(self._clear)
        lv = QVBoxLayout()
        lv.addWidget(self.logo, 0, Qt.AlignHCenter)
        lv.addWidget(pick)
        lv.addWidget(clear)
        lv.addStretch()

        save = variant(QPushButton("💾  Kaydet"), "primary")
        save.setMinimumHeight(44)
        save.setMinimumWidth(180)
        save.clicked.connect(self.save)
        self.admins_note = note = QLabel()
        note.setWordWrap(True)
        note.setObjectName("subtitle")

        card = QGroupBox("Okul")
        h = QHBoxLayout(card)
        h.setContentsMargins(20, 22, 20, 18)
        h.setSpacing(28)
        h.addLayout(lv)
        right = QVBoxLayout()
        right.addLayout(form)
        right.addWidget(note)
        right.addStretch()
        h.addLayout(right, 1)
        card.setMaximumWidth(900)
        foot = QHBoxLayout()
        foot.addStretch()
        foot.addWidget(save)
        wrap = QVBoxLayout()
        wrap.addWidget(card)
        wrap.addLayout(foot)
        center = QHBoxLayout()
        center.addStretch(1)
        center.addLayout(wrap, 6)
        center.addStretch(1)
        v = QVBoxLayout(self)
        v.setContentsMargins(0, 0, 0, 0)
        v.addWidget(page_header("🏫  Okul Bilgileri",
                                "Bu bilgiler yazdırılan çizelgelerde ve programın sol üst "
                                "köşesinde görünür."))
        v.addLayout(center)
        v.addStretch()
        self.refresh()

    def refresh(self):
        for key, e in self.fields.items():
            e.setText(self.db.get_setting(key, ""))
        self.new_logo, self.remove_logo = None, False
        self._show_logo(logo_path() if os.path.exists(logo_path()) else None)
        self.admins_note.setText(
            "Okul müdürü ve müdür yardımcıları, ders programı yüklendikten sonra "
            "<b>Ders Programı</b> sayfasının altındaki <b>İdareciler</b> bölümünden listeden "
            "seçilir. Çıktılardaki imza yerlerine bu seçim yazılır.")

    def _show_logo(self, path):
        pm = QPixmap(path) if path else QPixmap()
        if pm.isNull():
            self.logo.setPixmap(QPixmap())
            self.logo.setText("Logo yok")
        else:
            self.logo.setPixmap(pm.scaled(120, 120, Qt.KeepAspectRatio, Qt.SmoothTransformation))

    def _pick(self):
        path, _ = QFileDialog.getOpenFileName(self, "Logo seç", "",
                                              "Resim (*.png *.jpg *.jpeg *.bmp *.gif)")
        if path:
            self.new_logo, self.remove_logo = path, False
            self._show_logo(path)

    def _clear(self):
        self.new_logo, self.remove_logo = None, True
        self._show_logo(None)

    def save(self):
        if not self.fields["school"].text().strip():
            _info(self, "Lütfen okul adını yazın.")
            return False
        first = not self.db.get_setting("school")
        if self.new_logo:
            img = QImage(self.new_logo)
            if img.isNull():
                _info(self, "Seçilen resim açılamadı.")
                return False
            round_logo(img).save(logo_path())
        elif self.remove_logo and os.path.exists(logo_path()):
            os.remove(logo_path())
        for key, e in self.fields.items():
            self.db.set_setting(key, e.text().strip())
        self.refresh()
        self.win.school_saved(first)
        return True


class DutyEditDialog(QDialog):
    """Nöbet çizelgesini kaydetmeden önce gösterir/düzenler."""

    def __init__(self, parent, entries, teachers, month, unmatched=(), warnings=()):
        super().__init__(parent)
        self.setWindowTitle("Nöbet çizelgesi — kontrol edip kaydedin")
        self.resize(900, 650)
        self.teachers = sorted(teachers, key=normalize)
        self.matcher = NameMatcher(self.teachers)

        self.month = MonthPicker(month)

        self.table = QTableWidget(0, 4)
        self.table.setHorizontalHeaderLabels(["Gün", "Tarih (isteğe bağlı)", "Öğretmen", "Nöbet yeri"])
        align(self.table)
        _stretch(self.table, 2, 3)
        for e in sorted(entries, key=lambda e: (e.date or "", e.weekday, e.place, e.teacher)):
            self._add_row(e)

        add = QPushButton("Satır ekle")
        add.clicked.connect(lambda: self._add_row(DutyEntry(0, "", "")))
        rem = QPushButton("Seçili satırları sil")
        rem.clicked.connect(self._remove)
        hb = QHBoxLayout()
        hb.addWidget(add)
        hb.addWidget(rem)
        hb.addStretch()

        msgs = list(warnings)
        if unmatched:
            msgs.append("Ders programında bulunamayan isimler (eklenmedi): " + ", ".join(unmatched))
        info = QLabel("\n".join(msgs) or "Tüm isimler ders programıyla eşleşti.")
        info.setWordWrap(True)
        if msgs:
            info.setStyleSheet("color:#a05a00")

        buttons = QDialogButtonBox(QDialogButtonBox.Save | QDialogButtonBox.Cancel)
        buttons.button(QDialogButtonBox.Save).setText("Kaydet")
        buttons.button(QDialogButtonBox.Cancel).setText("Vazgeç")
        buttons.accepted.connect(self._accept)
        buttons.rejected.connect(self.reject)

        form = QFormLayout()
        form.addRow("Çizelge şu aydan itibaren geçerli:", self.month)
        v = QVBoxLayout(self)
        v.addLayout(form)
        v.addWidget(QLabel("Tarih sütunu boşsa çizelge her hafta o gün için geçerlidir."))
        v.addWidget(self.table, 1)
        v.addLayout(hb)
        v.addWidget(info)
        v.addWidget(buttons)

    def _add_row(self, e):
        r = self.table.rowCount()
        self.table.insertRow(r)
        day = QComboBox()
        day.addItems(DAYS)
        day.setCurrentIndex(e.weekday)
        self.table.setCellWidget(r, 0, day)
        self.table.setItem(r, 1, _item(e.date and _qdate(e.date).toString("dd.MM.yyyy"), True))
        who = QComboBox()
        who.setEditable(True)
        who.addItems(self.teachers)
        who.setCurrentText(e.teacher)
        comp = QCompleter(self.teachers)
        comp.setCaseSensitivity(Qt.CaseInsensitive)
        comp.setFilterMode(Qt.MatchContains)
        who.setCompleter(comp)
        self.table.setCellWidget(r, 2, who)
        self.table.setItem(r, 3, _item(e.place, True))
        if e.teacher and e.teacher not in self.teachers:
            who.setStyleSheet("background:#ffd6d6")
            who.setToolTip("Ders programında yok: bu öğretmenin hiç dersi yokmuş gibi "
                           "davranılır. Yanlış okunduysa düzeltin ya da satırı silin.")

    def _remove(self):
        for r in sorted({i.row() for i in self.table.selectedIndexes()}, reverse=True):
            self.table.removeRow(r)

    def _accept(self):
        bad = []
        for r in range(self.table.rowCount()):
            name = self.table.cellWidget(r, 2).currentText().strip()
            if name and name not in self.teachers and not self.matcher.match(name):
                bad.append(name)
        if bad and not _ask(self, "Şu isimler ders programında yok:\n" + "\n".join(bad) +
                            "\n\nYine de kaydedilsin mi? (Bu öğretmenlerin dersi olmadığı varsayılır)"):
            return
        self.accept()

    def entries(self):
        out = []
        for r in range(self.table.rowCount()):
            name = self.table.cellWidget(r, 2).currentText().strip()
            if not name:
                continue
            name = name if name in self.teachers else (self.matcher.match(name) or clean_name(name))
            wd = self.table.cellWidget(r, 0).currentIndex()
            raw = self.table.item(r, 1).text().strip() if self.table.item(r, 1) else ""
            date = None
            if raw:
                q = QDate.fromString(raw, "dd.MM.yyyy")
                if q.isValid():
                    date = _iso(q)
                    wd = min(q.dayOfWeek() - 1, 4)
            place = self.table.item(r, 3).text().strip() if self.table.item(r, 3) else ""
            out.append(DutyEntry(wd, name, place, date))
        return out

    def month_key(self):
        return self.month.key()


# ---------------------------------------------------------------- sekmeler

class DailyTab(QWidget):
    """Kademeli görevlendirme: 1) tarih + gelmeyenler  2) boş dersler + nöbetçiler
    3) sonuç (değiştir, yazdır). Her adım tam sayfadır."""

    def __init__(self, win):
        super().__init__()
        self.win, self.db = win, win.db
        self.plan = None
        self.approved = False
        self.schedule = {}
        self.duty = []
        self.admins = set()
        self.row_index = {}
        self.day_len = 8
        self.school_hours = 7

        self.stepbar = theme.StepBar(["Gelmeyen öğretmenleri seçin", "Boş dersler ve nöbetçiler",
                                      "Görevlendirme"])
        self.steps = QStackedWidget()

        # ---------- Adım 1: tarih ve o gün dersi olan öğretmenler
        self.date = QDateEdit(_qdate(_next_school_day(dt.date.today()).isoformat()))
        self.date.setCalendarPopup(True)
        self.date.setDisplayFormat("d MMMM yyyy, dddd")
        self.date.setMinimumHeight(40)
        self.date.setMinimumWidth(280)
        self.date.setStyleSheet("font-size:12pt; font-weight:600")
        self.date.dateChanged.connect(self.refresh)
        prev_b, next_b = QPushButton("◀"), QPushButton("▶")
        for btn in (prev_b, next_b):
            btn.setFixedSize(40, 40)
            btn.setStyleSheet("padding:0px; font-size:11pt")
            btn.setCursor(Qt.PointingHandCursor)
        prev_b.clicked.connect(lambda: self._shift(-1))
        next_b.clicked.connect(lambda: self._shift(1))
        self.filter = QLineEdit()
        self.filter.setPlaceholderText("🔍  Öğretmen ara…")
        self.filter.setMinimumHeight(38)
        self.filter.textChanged.connect(self._filter)
        top = QHBoxLayout()
        top.addWidget(QLabel("<b>Tarih</b>"))
        top.addWidget(prev_b)
        top.addWidget(self.date)
        top.addWidget(next_b)
        top.addSpacing(16)
        top.addWidget(self.filter, 1)
        self.status = QLabel()
        self.status.setWordWrap(True)
        self.saved_bar = QFrame()
        self.saved_bar.setObjectName("hint")
        sb = QHBoxLayout(self.saved_bar)
        sb.setContentsMargins(14, 8, 8, 8)
        sb.addWidget(QLabel("✅ Bu gün için kayıtlı bir görevlendirme var."), 1)
        show_saved = variant(QPushButton("Görevlendirmeyi göster"), "primary")
        show_saved.clicked.connect(lambda: self.go(2))
        sb.addWidget(show_saved)
        self.people = QTableWidget(0, 4)
        self.people.setHorizontalHeaderLabels(["Öğretmen", "Bugünkü ders saatleri", "Ders",
                                               "Durum"])
        _stretch(self.people, 0)
        align(self.people, 2)
        self.people.setItemDelegateForColumn(1, HourStripDelegate(self.people))
        for c, w in ((1, 260), (2, 64), (3, 190)):
            self.people.horizontalHeader().setSectionResizeMode(c, QHeaderView.Fixed)
            self.people.setColumnWidth(c, w)
        self.people.setShowGrid(False)
        self.people.verticalHeader().hide()
        self.people.verticalHeader().setDefaultSectionSize(54)
        self.people.setSelectionMode(QAbstractItemView.NoSelection)
        self.people.setFocusPolicy(Qt.NoFocus)
        self.people.setMinimumHeight(170)
        self.next1 = variant(QPushButton("Devam  →"), "primary")
        self.next1.setMinimumHeight(46)
        self.next1.setStyleSheet("font-size:12pt")
        self.next1.clicked.connect(lambda: self.go(1))
        self.picked = QLabel()
        self.picked.setObjectName("subtitle")
        self.picked.setWordWrap(True)
        f1 = QHBoxLayout()
        f1.addWidget(self.picked, 1)
        f1.addWidget(self.next1)
        p1 = QWidget()
        v1 = QVBoxLayout(p1)
        v1.setContentsMargins(0, 0, 0, 0)
        v1.addLayout(top)
        v1.addWidget(self.status)
        v1.addWidget(self.saved_bar)
        hint = QLabel("Gelmeyen öğretmenin satırında <b>Durum</b>u <b>Raporlu, İzinli, Görevli…</b> "
                      "yapın, sonra <b>Devam</b>'a basın.")
        hint.setObjectName("subtitle")
        v1.addWidget(hint)
        v1.addWidget(self.people, 1)
        v1.addLayout(f1)

        # ---------- Adım 2: boş kalan dersler ve nöbetçiler
        self.gaps = QTableWidget(0, 3)
        self.gaps.setHorizontalHeaderLabels(["Saat", "Sınıf", "Ders"])
        self.gaps.cellClicked.connect(self._edit_gap)
        self.gap_rows = {}
        align(self.gaps, 0, 1)
        self.dutyt = QTableWidget(0, 4)
        self.dutyt.setHorizontalHeaderLabels(["Sıra", "Nöbetçi öğretmen", "Boş kalan saatlerde",
                                              "Girdiği"])
        for c, tip in ((0, "Adalet sırası: 1 numara ilk görevlendirilecek kişidir (bu yıl en az "
                           "boş derse giren)"),
                       (3, "Bu eğitim yılında girdiği boş ders sayısı (onaylanmış "
                           "görevlendirmeler); parantez içi bu ay")):
            self.dutyt.horizontalHeaderItem(c).setToolTip(tip)
        for t, stretch_col in ((self.gaps, 2), (self.dutyt, 1)):
            _stretch(t, stretch_col)
            t.setMinimumHeight(150)
            t.setShowGrid(False)
            t.verticalHeader().hide()
            t.verticalHeader().setDefaultSectionSize(40)
            t.setSelectionMode(QAbstractItemView.NoSelection)
            t.setFocusPolicy(Qt.NoFocus)
        align(self.gaps, 0, 1)
        align(self.dutyt, 0, 3)
        self.dutyt.verticalHeader().setDefaultSectionSize(54)
        self.dutyt.setItemDelegateForColumn(2, HourStripDelegate(self.dutyt))
        for c, w in ((0, 48), (2, 120), (3, 70)):
            self.dutyt.horizontalHeader().setSectionResizeMode(c, QHeaderView.Fixed)
            self.dutyt.setColumnWidth(c, w)
        for c, w in ((0, 72), (1, 64)):
            self.gaps.horizontalHeader().setSectionResizeMode(c, QHeaderView.Fixed)
            self.gaps.setColumnWidth(c, w)
        legend = QLabel("<span style='color:#15803d'><b>boş</b></span> = o saatte derse girebilir"
                        " &nbsp;·&nbsp; <span style='color:#b91c1c'><b>dolu</b></span> = o saatte "
                        "kendi dersi var")
        legend.setWordWrap(True)
        legend.setObjectName("subtitle")
        self.gaps_box = QGroupBox("Boş kalan dersler")
        gl = QVBoxLayout(self.gaps_box)
        gl.setContentsMargins(12, 18, 12, 12)
        gl.addWidget(self.gaps)
        self.duty_box = QGroupBox("Bugünün nöbetçileri")
        dl = QVBoxLayout(self.duty_box)
        dl.setContentsMargins(12, 18, 12, 12)
        dl.addWidget(self.dutyt, 1)
        dl.addWidget(legend)

        # Joker öğretmenler: o gün okulda olan, elle eklenen öğretmenler
        add_j = variant(QPushButton("＋  Joker öğretmen ekle"), "primary")
        add_j.setMinimumHeight(38)
        add_j.clicked.connect(self._add_joker)
        self.joker_msg = QLabel()
        self.joker_msg.setObjectName("okMsg")
        self.chips = QHBoxLayout()
        self.chips.setSpacing(6)
        self.joker_box = QFrame()
        self.joker_box.setObjectName("jokerCard")
        jv = QVBoxLayout(self.joker_box)
        jv.setContentsMargins(16, 10, 12, 10)
        jv.setSpacing(6)
        row1 = QHBoxLayout()
        self.joker_title = QLabel("<b style='font-size:11pt'>🃏 Joker öğretmenler</b>")
        jsub = QLabel("Listede olmayan (ör. dışarıdan gelen) ya da bugün dersi olmayan öğretmenler."
                      " İlk önce bunlar görevlendirilir (günde 7 derse kadar).")
        jsub.setObjectName("subtitle")
        jsub.setWordWrap(True)
        row1.addWidget(self.joker_title)
        row1.addWidget(jsub, 1)
        row1.addWidget(add_j)
        row2 = QHBoxLayout()
        row2.addLayout(self.chips)
        row2.addStretch()
        row2.addWidget(self.joker_msg)
        jv.addLayout(row1)
        jv.addLayout(row2)
        g2 = QHBoxLayout()
        g2.setSpacing(14)
        g2.addWidget(self.gaps_box, 3)
        g2.addWidget(self.duty_box, 7)
        back2 = QPushButton("←  Geri")
        back2.setMinimumHeight(46)
        back2.clicked.connect(lambda: self.go(0))
        self.run_btn = variant(QPushButton("Görevlendir  ▶"), "primary")
        self.run_btn.setMinimumHeight(46)
        self.run_btn.setStyleSheet("font-size:12pt")
        self.run_btn.clicked.connect(lambda: self.run())
        f2 = QHBoxLayout()
        f2.addWidget(back2)
        info2 = QLabel("Sıra: 1) joker öğretmenler (günde 7 derse kadar), 2) o saatte boş "
                       "nöbetçiler (bu yıl en az derse giren önce), 3) o saatte dersi olmayan "
                       "diğer öğretmenler.")
        info2.setObjectName("subtitle")
        info2.setWordWrap(True)
        f2.addWidget(info2, 1)
        f2.addWidget(self.run_btn)
        p2 = QWidget()
        v2 = QVBoxLayout(p2)
        v2.setContentsMargins(0, 0, 0, 0)
        v2.addWidget(self.joker_box)
        v2.addLayout(g2, 1)
        v2.addLayout(f2)

        # ---------- Adım 3: sonuç
        self.result = QTableWidget(0, 4)
        self.result.setHorizontalHeaderLabels(["Saat", "Sınıf", "Ders", "Derse girecek öğretmen"])
        align(self.result, 0, 1)
        _stretch(self.result, 3)
        for c, w in ((0, 80), (1, 70), (2, 200)):
            self.result.horizontalHeader().setSectionResizeMode(c, QHeaderView.Fixed)
            self.result.setColumnWidth(c, w)
        self.result.setAlternatingRowColors(False)
        self.result.verticalHeader().hide()
        self.result.verticalHeader().setDefaultSectionSize(42)
        self.result.setSelectionMode(QAbstractItemView.NoSelection)
        self.result.setFocusPolicy(Qt.NoFocus)
        self.result.setMinimumHeight(170)
        self.summary = QLabel()
        self.summary.setObjectName("subtitle")
        self.summary.setWordWrap(True)
        back3 = QPushButton("←  Geri")
        back3.clicked.connect(lambda: self.go(1))
        prt = variant(QPushButton("🖨  Yazdır"), "primary")
        prt.clicked.connect(self.print_)
        xlsx = QPushButton("Excel")
        xlsx.clicked.connect(self.to_excel)
        new = QPushButton("Yeni görevlendirme")
        new.clicked.connect(lambda: self.go(0))
        self.approve_btn = variant(QPushButton("✅  Onayla ve istatistiğe işle"), "success")
        self.approve_btn.clicked.connect(self.approve)
        delete = variant(QPushButton("Bu günün görevlendirmesini sil"), "danger")
        delete.clicked.connect(self.delete)
        f3 = QHBoxLayout()
        for btn in (back3, self.approve_btn, prt, xlsx, new):
            btn.setMinimumHeight(40)
            f3.addWidget(btn)
        f3.addStretch()
        f3.addWidget(delete)
        hint3 = QLabel("Değiştirmek istediğiniz satırda <b>Derse girecek öğretmen</b> kutusundan "
                       "başka birini seçebilirsiniz. Her şey uygunsa <b>Onayla</b>'ya basın; görevlendirme "
                       "ancak onaylanınca istatistiğe işlenir.")
        hint3.setObjectName("subtitle")
        hint3.setWordWrap(True)
        p3 = QWidget()
        v3 = QVBoxLayout(p3)
        v3.setContentsMargins(0, 0, 0, 0)
        v3.addWidget(hint3)
        self.abs_list = QListWidget()
        self.abs_list.setObjectName("absentList")
        self.abs_list.setFixedWidth(250)
        self.abs_list.setCursor(Qt.PointingHandCursor)
        self.abs_list.setSelectionMode(QAbstractItemView.SingleSelection)
        self.abs_list.setDragEnabled(False)
        self.abs_list.currentRowChanged.connect(lambda _r: self.show_plan(keep_list=True))
        abs_box = QVBoxLayout()
        abs_lbl = QLabel("<b>Gelmeyen öğretmenler</b>")
        abs_box.addWidget(abs_lbl)
        abs_box.addWidget(self.abs_list, 1)
        body3w = QWidget()
        body3 = QHBoxLayout(body3w)
        body3.setContentsMargins(0, 0, 0, 0)
        body3.setSpacing(12)
        body3.addLayout(abs_box)
        body3.addWidget(self.result, 1)
        # İkinci görünüm: görevlendirilen öğretmenlerin o günkü programı (tek satır)
        self.progs = QTableWidget(0, 1)
        self.progs.verticalHeader().hide()
        self.progs.setSelectionMode(QAbstractItemView.NoSelection)
        self.progs.setFocusPolicy(Qt.NoFocus)
        self.progs.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.progs.setWordWrap(True)
        self.view3 = QStackedWidget()
        self.view3.addWidget(body3w)
        self.view3.addWidget(self.progs)
        seg = QHBoxLayout()
        seg.setSpacing(0)
        self.seg_btns = []
        for i, text in enumerate(("📋  Gelmeyen öğretmenlere göre",
                                  "👤  Görevlendirilenlerin günlük programı")):
            b = QPushButton(text)
            b.setCheckable(True)
            b.setObjectName("segment")
            b.setCursor(Qt.PointingHandCursor)
            b.clicked.connect(lambda _c=False, k=i: self._view3(k))
            self.seg_btns.append(b)
            seg.addWidget(b)
        seg.addStretch()
        self.seg_btns[0].setChecked(True)
        v3.addLayout(seg)
        v3.addWidget(self.view3, 1)
        v3.addWidget(self.summary)
        v3.addLayout(f3)

        for page in (p1, p2, p3):
            self.steps.addWidget(page)

        v = QVBoxLayout(self)
        v.setContentsMargins(0, 0, 0, 0)
        v.setSpacing(10)
        v.addWidget(page_header("📝  Görevlendirme"))
        v.addWidget(self._build_setup())
        v.addWidget(self.stepbar)
        v.addWidget(self.steps, 1)
        keep_still(self.people, self.gaps, self.dutyt, self.result, self.progs, self.abs_list)
        self.refresh()

    def go(self, index):
        if index >= 1 and not self.absences():
            _info(self, "Önce gelmeyen öğretmenin durumunu seçin (Raporlu, İzinli, Görevli…).")
            return
        if index == 2 and not self.plan:
            self.run()
            return
        if index == 1:
            self._sections()
        self.steps.setCurrentIndex(index)
        self.stepbar.set(index)

    def _build_setup(self):
        """Ders programı veya nöbet çizelgesi yoksa gösterilen 'Başlarken' bölümü."""
        self.setup = QFrame()
        self.setup.setObjectName("hint")
        self.step0 = variant(QPushButton(), "step")
        self.step1 = variant(QPushButton(), "step")
        self.step2 = variant(QPushButton(), "step")
        self.step0.clicked.connect(self.win.edit_school)
        self.step1.clicked.connect(lambda: (self.win.sched.load_file(), self.refresh()))
        self.step2.clicked.connect(lambda: (self.win.duty.load_file(), self.refresh()))
        manual = QLabel("Belgeniz okunamazsa: <b>Ders Programı</b> bölümünde <i>Öğretmen ekle</i> "
                        "ile, <b>Nöbet Çizelgesi</b> bölümünde <i>Elle oluştur / düzenle</i> "
                        "ile elle girebilirsiniz.")
        manual.setWordWrap(True)
        title = QLabel("<b style='font-size:12pt'>👋 Hoş geldiniz!</b> — Programı kullanmaya "
                       "başlamak için bu üç adımı sırayla tamamlayın:")
        lay = QVBoxLayout(self.setup)
        lay.setContentsMargins(16, 12, 16, 12)
        lay.addWidget(title)
        h = QHBoxLayout()
        for b in (self.step0, self.step1, self.step2):
            h.addWidget(b)
        lay.addLayout(h)
        lay.addWidget(manual)
        return self.setup

    def _update_setup(self, has_schedule, has_duty):
        has_school = bool(self.db.get_setting("school", ""))
        self.step0.setText(("✅  " if has_school else "1️⃣  ") + "Okul bilgilerini girin\n"
                           "      okul adı, il/ilçe, logo")
        self.step1.setText(("✅  " if has_schedule else "2️⃣  ") + "Ders programını yükleyin\n"
                           "      PDF veya Excel")
        self.step2.setText(("✅  " if has_duty else "3️⃣  ") + "Nöbet çizelgesini yükleyin\n"
                           "      bu ayın çizelgesi")
        self.step2.setEnabled(has_schedule)
        self.setup.setVisible(not (has_school and has_schedule and has_duty))

    def _shift(self, n):
        d = self.date.date().addDays(n)
        while d.dayOfWeek() > 5:
            d = d.addDays(n)
        self.date.setDate(d)

    def date_iso(self):
        return _iso(self.date.date())

    def refresh(self):
        date = self.date_iso()
        wd = self.date.date().dayOfWeek() - 1
        self.plan = None
        self._update_setup(bool(self.db.schedule_versions()), bool(self.db.duty_months()))
        self.win.update_brand()
        self.people.setRowCount(0)
        self.saved_bar.hide()
        self.steps.setCurrentIndex(0)
        self.stepbar.set(0)
        if wd > 4:
            self.schedule, self.duty = {}, []
            self.status.setText("Hafta sonu — okul yok.")
            self._sections()
            return
        vid = self.db.schedule_version_for(date)
        self.schedule = self.db.schedule(vid) if vid else {}
        duty, month = self.db.duty_for(date)
        self.admins = self.db.admins()
        self.duty = [(t, p) for t, p in duty if t not in self.admins and not is_admin_place(p)]
        msgs = []
        if not self.schedule:
            msgs.append("<span style='color:#dc2626'>Ders programı yüklenmedi.</span>")
        if not month:
            msgs.append("<span style='color:#dc2626'>Nöbet çizelgesi yüklenmedi.</span>")
        elif month != date[:7]:
            msgs.append(f"<span style='color:#dc2626'><b>{export.month_tr(date[:7])} nöbet "
                        f"çizelgesi yüklenmedi!</b> Şimdilik {export.month_tr(month)} çizelgesi "
                        "kullanılıyor — <b>Nöbet Çizelgesi</b> bölümünden bu ayınkini "
                        "yükleyin.</span>")
        # İdareciler listelenmez: onlara zaten boş ders görevlendirmesi yapılmaz
        teaching = sorted((t for t in teachers_with_lessons(self.schedule, wd)
                           if t not in self.admins), key=normalize)
        self.day_len = max([8] + [h for slots in self.schedule.values() for (d, h) in slots
                                  if d == wd])
        # Okulun o günkü gerçek ders saati sayısı (boş saat hesabı için)
        self.school_hours = max([h for slots in self.schedule.values() for (d, h) in slots
                                 if d == wd] or [7])
        if teaching:
            msgs.append(f"<span style='color:{theme.MUTED}'>Bugün {len(teaching)} öğretmenin dersi "
                        f"var, {len({t for t, _ in self.duty})} nöbetçi.</span>")
        self.status.setText("<br>".join(msgs))

        saved = self.db.assignments(date)
        saved_abs = self.db.absences(date)
        self.people.setRowCount(len(teaching))
        for r, n in enumerate(teaching):
            k = len(self._lessons_today(n))
            it = _item(f"{n}")
            self.people.setItem(r, 0, it)
            strip = _item("")
            strip.setData(Qt.UserRole, {"max": self.day_len, "lessons": set(self._lessons_today(n))})
            strip.setToolTip(", ".join(map(str, self._lessons_today(n))) + ". saatlerde dersi var")
            self.people.setItem(r, 1, strip)
            self.people.setItem(r, 2, _item(k))
            combo = ChoiceButton(STATUSES)
            if n in saved_abs:
                reason = saved_abs[n]
                if reason not in STATUSES:
                    combo.insertItem(combo.count() - 1, reason)
                combo.setCurrentText(reason)
            combo.currentTextChanged.connect(lambda _t, row=r: self._status_changed(row))
            self.people.setCellWidget(r, 3, combo)
            self._mark(r)
        self._filter(self.filter.text())
        if saved_abs:
            self.run(saved=saved)
            self.saved_bar.show()
            self.steps.setCurrentIndex(0)
            self.stepbar.set(0)
        self._update_picked()

    def _lessons_today(self, teacher):
        wd = self.date.date().dayOfWeek() - 1
        return sorted(h for (d, h) in self.schedule.get(teacher, {}) if d == wd)

    def _status_changed(self, row):
        combo = self.people.cellWidget(row, 3)
        if combo.currentText() == OTHER:
            name = self.people.item(row, 0).text()
            dlg = QInputDialog(self)
            dlg.setWindowTitle("Mazeret")
            dlg.setLabelText(f"{name} için mazereti yazın:")
            dlg.setOkButtonText("Tamam")
            dlg.setCancelButtonText("Vazgeç")
            ok = dlg.exec() == QDialog.Accepted
            text = dlg.textValue()
            dlg.deleteLater()
            text = text.strip()
            combo.blockSignals(True)
            if ok and text:
                if combo.findText(text) < 0:
                    combo.insertItem(combo.count() - 1, text)
                combo.setCurrentText(text)
            else:
                combo.setCurrentText("Okulda")
            combo.blockSignals(False)
        self._mark(row)
        self.plan = None  # seçim değişti: görevlendirme yeniden yapılmalı
        if not self.absences() and self.db.absences(self.date_iso()):
            self.db.delete_day(self.date_iso())
            self.win.write_excel()
        self._update_picked()

    def _mark(self, row):
        combo = self.people.cellWidget(row, 3)
        absent = combo.currentText() != "Okulda"
        combo.setProperty("state", "absent" if absent else "present")
        combo.style().unpolish(combo)
        combo.style().polish(combo)
        it = self.people.item(row, 0)
        f = it.font()
        f.setBold(absent)
        it.setFont(f)
        it.setForeground(QColor("#dc2626") if absent else QColor(theme.TEXT_COLOR))

    def _filter(self, text):
        key = normalize(text)
        for r in range(self.people.rowCount()):
            self.people.setRowHidden(r, bool(key) and key not in normalize(self.people.item(r, 0).text()))

    def absences(self):
        out = {}
        for r in range(self.people.rowCount()):
            combo = self.people.cellWidget(r, 3)
            s = combo.currentText() if combo else "Okulda"
            if s != "Okulda":
                out[self.people.item(r, 0).text()] = s
        return out

    def _update_picked(self):
        absent = self.absences()
        if absent:
            n = sum(len(self._lessons_today(t)) for t in absent)
            self.picked.setText(f"<b>{len(absent)}</b> gelmeyen öğretmen seçildi: "
                                + ", ".join(f"{t} ({r})" for t, r in absent.items())
                                + f" — <b>{n}</b> ders boş kalıyor.")
        else:
            self.picked.setText("Henüz gelmeyen öğretmen seçilmedi.")
        self.next1.setEnabled(bool(absent))

    def _sections(self):
        """2. adımın tablolarını doldurur."""
        absent = self.absences()
        if not absent:
            return
        wd = self.date.date().dayOfWeek() - 1
        # Gelmeyenlerin dersleri: öğretmen başlığı + Saat | Sınıf | Ders
        self.gaps.clearSpans()
        self.gaps.clearContents()
        self.gaps.setRowCount(0)
        self.gap_rows = {}
        need = set()
        total = 0
        for t, reason in absent.items():
            r = self.gaps.rowCount()
            self.gaps.insertRow(r)
            k = len(self._lessons_today(t))
            head = _item(f"{t}  ·  {reason}  ·  {k} saat ders, "
                         f"{max(0, self.school_hours - k)} saat boş")
            f = head.font()
            f.setBold(True)
            head.setFont(f)
            head.setBackground(BLUE)
            self.gaps.setItem(r, 0, head)
            self.gaps.setSpan(r, 0, 1, 3)
            for h in self._lessons_today(t):
                need.add(h)
                total += 1
                sinif, ders = split_info(self.schedule[t][(wd, h)])
                r = self.gaps.rowCount()
                self.gaps.insertRow(r)
                self.gap_rows[r] = (t, h)
                self.gaps.setItem(r, 0, _item(f"{h}. ders"))
                destek = "destek" in normalize(ders or "")
                c_it = _item(sinif or ("—" if destek else "✎ seçin"))
                d_it = _item(lesson_label(ders) or "✎ yazın")
                for it, empty in ((c_it, not sinif and not destek), (d_it, not ders)):
                    it.setToolTip("Düzeltmek için tıklayın")
                    if empty:
                        it.setBackground(YELLOW)
                        f = it.font()
                        f.setBold(True)
                        it.setFont(f)
                self.gaps.setItem(r, 1, c_it)
                self.gaps.setItem(r, 2, d_it)
        # Nöbetçiler: yalnızca boş kalan derslerin saatlerinden hangilerinde boş oldukları
        date = self.date_iso()
        year = self.db.year_counts(date, exclude_date=date)
        month = self.db.counts(date[:7], exclude_date=date)
        # Adalet sırası: bu yıl en az boş derse giren önce (eşitse bu ay, sonra ad)
        names = sorted({t for t, _ in self.duty},
                       key=lambda t: (year.get(t, (0,))[0], month.get(t, (0,))[0], normalize(t)))
        self.dutyt.clearContents()
        self.dutyt.setRowCount(len(names))
        for r, t in enumerate(names):
            rank = _item(f"{r + 1}")
            name = _item(t)
            if r == 0:
                f = name.font()
                f.setBold(True)
                name.setFont(f)
            strip = _item("")
            if t in absent:
                strip.setData(Qt.UserRole, {"max": self.day_len, "need": set(), "busy": set()})
                name.setText(f"{t}  (bugün gelmedi)")
                name.setForeground(QColor("#dc2626"))
            else:
                busy = set(self._lessons_today(t))
                strip.setData(Qt.UserRole, {"max": self.day_len, "need": set(need),
                                            "busy": busy & set(need)})
                fit = sorted(h for h in need if h not in busy)
                strip.setToolTip(("Girebileceği saatler: " + ", ".join(map(str, fit)))
                                 if fit else "Boş kalan saatlerin hepsinde kendi dersi var")
            count = _item(f"{year.get(t, (0,))[0]}  ({month.get(t, (0,))[0]})")
            for c, it in enumerate((rank, name, strip, count)):
                self.dutyt.setItem(r, c, it)
        self.dutyt.setColumnWidth(2, max(120, 20 + len(need) * (HourStripDelegate.BOX +
                                                                HourStripDelegate.GAP)))
        self._fill_jokers()
        missing = any(needs_class(self.schedule[t].get((wd, h), "") or "")
                      for t, h in self.gap_rows.values())
        self.gaps_box.setTitle(f"Boş kalan dersler — {total} ders" +
                               (" · sarıya tıklayıp düzeltin" if missing else ""))
        self.duty_box.setTitle(f"Bugünün nöbetçileri — {len(names)} kişi, adalet sırasıyla")

    def _fill_jokers(self):
        jokers = self.db.jokers(self.date_iso())
        while self.chips.count():
            w = self.chips.takeAt(0).widget()
            if w:
                w.hide()
                w.deleteLater()
        if not jokers:
            empty = QLabel("Henüz joker öğretmen eklenmedi.")
            empty.setObjectName("subtitle")
            self.chips.addWidget(empty)
        for t in jokers:
            chip = QPushButton(f"{t}   ✕")
            chip.setObjectName("chip")
            chip.setCursor(Qt.PointingHandCursor)
            chip.setToolTip("Listeden çıkarmak için tıklayın")
            chip.clicked.connect(lambda _c=False, who=t: self._remove_joker(who))
            self.chips.addWidget(chip)
        self.joker_title.setText(f"<b style='font-size:11pt'>🃏 Joker öğretmenler"
                                 f"{f' ({len(jokers)})' if jokers else ''}</b>")

    def _add_joker(self):
        """Joker seçme penceresi: önce listeden seçilir; listede yoksa yeni ad yazılıp eklenir."""
        date = self.date_iso()
        wd = self.date.date().dayOfWeek() - 1
        duty = {t for t, _ in self.duty}
        jokers = self.db.jokers(date)
        absent = self.absences()
        busy_today = set(teachers_with_lessons(self.schedule, wd))
        names = sorted((t for t in self.db.all_teachers()
                        if t not in duty and t not in jokers and t not in absent
                        and t not in self.admins),
                       key=lambda t: (t in busy_today, normalize(t)))

        dlg = QDialog(self)
        dlg.setWindowTitle("Joker öğretmen ekle")
        dlg.resize(520, 600)
        search = QLineEdit()
        search.setPlaceholderText("🔍  Listede ara…")
        search.setMinimumHeight(36)
        lst = QListWidget()
        lst.setObjectName("list")

        def add_row(t, checked=False, new=False):
            note = ("  — yeni eklendi" if new else
                    "  — bugün dersi yok" if t not in busy_today else
                    f"  — bugün {len(self._lessons_today(t))} dersi var")
            it = QListWidgetItem(t + note)
            it.setData(Qt.UserRole, t)
            # İşaret, satırın herhangi bir yerine tıklanınca değişir (kutu dahil, tek tıkla)
            it.setFlags(it.flags() & ~Qt.ItemIsUserCheckable)
            it.setCheckState(Qt.Checked if checked else Qt.Unchecked)
            return it

        for t in names:
            lst.addItem(add_row(t))
        lst.itemClicked.connect(lambda it: it.setCheckState(
            Qt.Unchecked if it.checkState() == Qt.Checked else Qt.Checked))

        def filt(text):
            key = normalize(text)
            for i in range(lst.count()):
                lst.item(i).setHidden(bool(key) and key not in normalize(lst.item(i).text()))
        search.textChanged.connect(filt)

        # Listede yoksa: yeni ad ekleme alanı
        new_name = QLineEdit()
        new_name.setPlaceholderText("Adı soyadı (ör. Ayşe Yılmaz)")
        new_name.setMinimumHeight(36)
        new_btn = QPushButton("＋  Listeye ekle")
        new_btn.setMinimumHeight(36)
        new_msg = QLabel()
        new_msg.setObjectName("okMsg")

        def add_new():
            name = clean_name(new_name.text())
            if not name:
                new_msg.setText("<span style='color:#b45309'>Önce adı soyadı yazın.</span>")
                return
            for i in range(lst.count()):
                if normalize(lst.item(i).data(Qt.UserRole)) == normalize(name):
                    lst.item(i).setCheckState(Qt.Checked)
                    lst.scrollToItem(lst.item(i))
                    new_msg.setText(f"ℹ {name} zaten listede, işaretlendi.")
                    new_name.clear()
                    return
            it = add_row(name, checked=True, new=True)
            lst.insertItem(0, it)
            lst.scrollToItem(it)
            search.clear()
            new_name.clear()
            new_msg.setText(f"✅ {name} listeye eklendi ve işaretlendi.")
        new_btn.clicked.connect(add_new)
        new_name.returnPressed.connect(add_new)

        bb = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        bb.button(QDialogButtonBox.Ok).setText("İşaretlileri joker yap")
        variant(bb.button(QDialogButtonBox.Ok), "primary")
        bb.button(QDialogButtonBox.Cancel).setText("Vazgeç")
        bb.accepted.connect(dlg.accept)
        bb.rejected.connect(dlg.reject)
        newbox = QGroupBox("Listede yok mu? Buradan ekleyin")
        nb = QVBoxLayout(newbox)
        nb.setContentsMargins(12, 18, 12, 10)
        nr = QHBoxLayout()
        nr.addWidget(new_name, 1)
        nr.addWidget(new_btn)
        nb.addLayout(nr)
        nb.addWidget(new_msg)
        v = QVBoxLayout(dlg)
        v.addWidget(page_header("🃏  Joker öğretmen ekle",
                                "1) Önce listeden bakın ve işaretleyin. 2) Listede yoksa aşağıya "
                                "adını yazıp ekleyin."))
        v.addWidget(search)
        v.addWidget(lst, 1)
        v.addWidget(newbox)
        v.addWidget(bb)
        search.setFocus()
        if not run_dialog(dlg):
            return
        picked = [lst.item(i).data(Qt.UserRole) for i in range(lst.count())
                  if lst.item(i).checkState() == Qt.Checked]
        if not picked:
            _info(self, "Hiç öğretmen işaretlenmedi.", "Joker öğretmen")
            return
        self.db.set_jokers(date, jokers + [t for t in picked if t not in jokers])
        self.plan = None
        self._fill_jokers()
        self.joker_msg.setText(f"✅ {len(picked)} joker öğretmen listeye eklendi.")
        _info(self, "✅ Joker öğretmen listesine eklendi:\n\n• " + "\n• ".join(picked),
              "Joker öğretmen")

    def _remove_joker(self, who):
        self.db.set_jokers(self.date_iso(), [t for t in self.db.jokers(self.date_iso())
                                             if t != who])
        self.joker_msg.setText(f"{who} joker listesinden çıkarıldı.")
        self.plan = None
        self._fill_jokers()

    def _edit_gap(self, row, col):
        """Sınıfı / dersi okunamayan hücreyi o anda düzeltme (ders programına da kaydedilir)."""
        if row not in self.gap_rows or col not in (1, 2):
            return
        t, h = self.gap_rows[row]
        wd = self.date.date().dayOfWeek() - 1
        sinif, ders = split_info(self.schedule[t].get((wd, h), "") or "")
        dlg = QInputDialog(self)
        dlg.setOkButtonText("Kaydet")
        dlg.setCancelButtonText("Vazgeç")
        if col == 1:
            classes = sorted({split_info(i or "")[0] for slots in self.schedule.values()
                              for i in slots.values()} - {""}, key=normalize)
            dlg.setWindowTitle("Sınıfı seçin")
            dlg.setLabelText(f"{t} — {h}. ders ({lesson_label(ders) or 'ders'}) hangi sınıfta?")
            dlg.setComboBoxItems(classes)
            dlg.setComboBoxEditable(True)
            dlg.setTextValue(sinif)
        else:
            dlg.setWindowTitle("Dersin adı")
            dlg.setLabelText(f"{t} — {h}. ders ({sinif or 'sınıf yok'}) hangi ders?")
            dlg.setTextValue(ders)
        ok = dlg.exec() == QDialog.Accepted
        value = clean_name(dlg.textValue()) if col == 2 else dlg.textValue().strip()
        dlg.deleteLater()
        if not ok:
            return
        if col == 1:
            sinif = value
        else:
            ders = value
        info = f"{sinif} {ders}".strip()
        self.schedule[t][(wd, h)] = info
        vid = self.db.schedule_version_for(self.date_iso())
        if vid:
            self.db.update_slot(vid, t, wd, h, info)
        self.plan = None
        self._sections()
        self.win.statusBar().showMessage(f"✅ {t} {h}. ders düzeltildi: {info} (ders programına "
                                         "da kaydedildi).", 6000)

    def run(self, saved=None):
        date = self.date_iso()
        if not self.absences():
            _info(self, "Önce soldaki listeden gelmeyen öğretmenin durumunu seçin "
                        "(Raporlu, İzinli, Görevli…).")
            return
        counts = self.db.counts(date[:7], exclude_date=date)
        totals = self.db.total_counts(exclude_date=date)
        cap = int(self.db.get_setting("max_per_day", "0") or 0)
        year = self.db.year_counts(date, exclude_date=date)
        self.plan = plan_day(date, self.schedule, self.duty, self.absences(), counts, totals, cap,
                             excluded=self.admins, year_counts=year,
                             jokers=self.db.jokers(date))
        if saved:
            by_key = {(h, t): (a, k, n) for h, t, _info_, a, k, n in saved}
            for a in self.plan.assignments:
                if (a.hour, a.absent) in by_key:
                    a.assigned, a.kind, a.note = by_key[(a.hour, a.absent)]
        self._sections()
        # Kayıtlı gün onaylı açılır; yeni görevlendirme onaylanana kadar taslaktır.
        self.approved = bool(saved)
        self.show_plan()
        self.steps.setCurrentIndex(2)
        self.stepbar.set(2)

    def _view3(self, k):
        for i, b in enumerate(self.seg_btns):
            b.setChecked(i == k)
        self.view3.setCurrentIndex(k)

    def _fill_programs(self):
        """Görevlendirilen öğretmenlerin o günkü programı: satır öğretmen, sütun ders saati."""
        n, rows = export.teacher_programs(self.plan, self.schedule)
        t = self.progs
        t.clear()
        t.setColumnCount(n + 1)
        t.setHorizontalHeaderLabels(["Öğretmen"] + [f"{h}. ders" for h in range(1, n + 1)])
        t.setRowCount(len(rows))
        h = t.horizontalHeader()
        h.setSectionResizeMode(0, QHeaderView.Fixed)
        t.setColumnWidth(0, 250)
        t.setWordWrap(False)
        for c in range(1, n + 1):
            h.setSectionResizeMode(c, QHeaderView.Stretch)
        for r, (name, cells) in enumerate(rows):
            k = sum(1 for kind, _ in cells.values() if kind == "task")
            it = _item(f"{name}  ({k})")
            it.setToolTip(f"{name}: bugün {k} ders görevlendirme")
            f = it.font()
            f.setBold(True)
            f.setPointSizeF(f.pointSizeF() + 1)
            it.setFont(f)
            it.setBackground(QColor(theme.ACCENT))
            it.setForeground(QColor("white"))
            t.setItem(r, 0, it)
            for hour in range(1, n + 1):
                kind, text = cells.get(hour, ("free", "—"))
                if kind == "task":
                    # Kısa gösterim: sınıf ve ders, altında kimin yerine (soyadıyla)
                    label, who = (text.split("\n") + [""])[:2]
                    who = who.strip("()").replace(" yerine", "")
                    cell = _item(f"{label}\n↳ {who.split()[-1] if who else ''}")
                else:
                    cell = _item(text)
                cell.setTextAlignment(Qt.AlignCenter)
                if kind == "task":
                    cell.setBackground(YELLOW)
                    cf = cell.font()
                    cf.setBold(True)
                    cell.setFont(cf)
                    cell.setToolTip("Görevlendirilen ders: " + text.replace("\n", " "))
                else:
                    cell.setForeground(QColor("#94a3b8" if kind == "free" else theme.MUTED))
                t.setItem(r, hour, cell)
        t.resizeRowsToContents()
        for r in range(t.rowCount()):
            t.setRowHeight(r, max(40, t.rowHeight(r) + 4))
        self.seg_btns[1].setText(f"👤  Görevlendirilenlerin günlük programı ({len(rows)})")

    def show_plan(self, keep_list=False):
        p = self.plan
        if p is None:
            return
        self._fill_programs()
        self.result.clearSpans()
        self.result.setRowCount(0)
        self.row_index = {}
        groups = {}
        for i, a in enumerate(p.assignments):
            groups.setdefault(a.absent, []).append(i)
        # Solda gelmeyen öğretmenler; tıklanan öğretmenin görevlendirmesi sağda açılır
        names = list(groups)
        if not keep_list:
            cur = self.abs_list.currentItem().data(Qt.UserRole) if self.abs_list.currentItem() \
                else None
            self.abs_list.blockSignals(True)
            self.abs_list.clear()
            for t in names:
                idxs = groups[t]
                missing = sum(1 for i in idxs if not p.assignments[i].assigned)
                it = QListWidgetItem(f"{t}\n{p.absences.get(t, '')} · {len(idxs)} saat ders, "
                                     f"{max(0, self.school_hours - len(idxs))} boş" +
                                     (f" · {missing} atanamadı" if missing else ""))
                it.setData(Qt.UserRole, t)
                self.abs_list.addItem(it)
            row = names.index(cur) if cur in names else 0
            self.abs_list.setCurrentRow(row)
            self.abs_list.blockSignals(False)
        sel = self.abs_list.currentItem().data(Qt.UserRole) if self.abs_list.currentItem() \
            else (names[0] if names else None)
        for absent, idxs in groups.items():
            if absent != sel:
                continue
            r = self.result.rowCount()
            self.result.insertRow(r)
            k = len(idxs)
            head = _item(f"{absent}   ·   {p.absences.get(absent, '')}   ·   {k} saat ders, "
                         f"{max(0, self.school_hours - k)} saat boş")
            f = head.font()
            f.setBold(True)
            head.setFont(f)
            head.setBackground(BLUE)
            self.result.setItem(r, 0, head)
            self.result.setSpan(r, 0, 1, 4)
            for i in idxs:
                a = p.assignments[i]
                r = self.result.rowCount()
                self.result.insertRow(r)
                self.row_index[r] = i
                sinif, ders = split_info(a.info)
                self.result.setItem(r, 0, _item(f"{a.hour}. ders"))
                self.result.setItem(r, 1, _item(sinif or "—"))
                self.result.setItem(r, 2, _item(lesson_label(ders) or a.info))
                combo = QComboBox()
                combo.addItem("— atanmadı —", None)
                opts = options_for(p, self.schedule, i)
                for t, lab in opts:
                    tag = {"nöbetçi": "", "idareci": "  (idareci)",
                           "joker": "  (joker)"}.get(lab, "  (nöbetçi değil)")
                    combo.addItem(t + tag, (t, lab))
                if a.assigned and a.assigned not in [t for t, _ in opts]:
                    combo.addItem(a.assigned, (a.assigned, "kayıtlı"))
                idx = next((k for k in range(combo.count())
                            if combo.itemData(k) and combo.itemData(k)[0] == a.assigned), 0)
                combo.setCurrentIndex(idx)
                combo.currentIndexChanged.connect(lambda _k, row=i, cb=combo: self._changed(row, cb))
                color = {"bos": YELLOW, "yok": RED, "joker": PURPLE}.get(a.kind)
                if color:
                    combo.setStyleSheet("font-weight:600")
                    for c in range(3):
                        self.result.item(r, c).setBackground(color)
                        self.result.item(r, c).setToolTip(a.note)
                self.result.setCellWidget(r, 3, combo)
        n = len(p.assignments)
        ok = sum(1 for a in p.assignments if a.kind == "nobetci")
        free = sum(1 for a in p.assignments if a.kind == "bos")
        joker = sum(1 for a in p.assignments if a.kind == "joker")
        none = sum(1 for a in p.assignments if not a.assigned)
        parts = [f"<b>{n}</b> ders", f"<b>{ok}</b> nöbetçiye verildi"]
        if joker:
            parts.append(f"<b style='color:#7e22ce'>{joker}</b> joker öğretmene verildi "
                         "(mor satırlar)")
        if free:
            parts.append(f"<b style='color:#d97706'>{free}</b> dersi olmayan diğer öğretmene "
                         "verildi (sarı satırlar)")
        if none:
            parts.append(f"<b style='color:#dc2626'>{none}</b> ders için boş öğretmen yok "
                         "(kırmızı satırlar)")
        if self.approved:
            parts.append("<b style='color:#16a34a'>✅ Onaylandı — istatistiğe işlendi</b>")
        else:
            parts.append("<b style='color:#d97706'>⏳ Onay bekliyor — henüz istatistiğe "
                         "işlenmedi</b>")
        self.approve_btn.setEnabled(not self.approved)
        self.approve_btn.setText("✅  Onaylandı" if self.approved
                                 else "✅  Onayla ve istatistiğe işle")
        self.summary.setText("  ·  ".join(parts))

    def _changed(self, row, combo):
        a = self.plan.assignments[row]
        data = combo.currentData()
        if not data:
            a.assigned, a.kind = None, "yok"
            a.note = "Elle boş bırakıldı."
        else:
            t, lab = data
            a.assigned = t
            if lab == "nöbetçi":
                a.kind, a.note = "nobetci", "Elle seçildi."
            elif lab == "idareci":
                a.kind, a.note = "bos", f"Elle seçildi; {t} idareci."
            elif lab == "joker":
                a.kind, a.note = "joker", f"Elle seçildi; {t} joker öğretmen."
            else:
                a.kind = "bos"
                a.note = f"⚠ Elle seçildi; {t} nöbetçi değil ({lab})."
        self.approved = False  # değişiklik yeniden onaylanmalı
        self.show_plan()

    def save(self):
        if not self.plan:
            _info(self, "Kaydedilecek görevlendirme yok.")
            return
        self.approve()

    def approve(self):
        if not self.plan:
            return
        self._store()
        self.approved = True
        self.show_plan()
        self.saved_bar.show()

    def _ensure_approved(self):
        """Yazdırma/Excel öncesi onaylanmamış taslak için onay ister."""
        if self.approved:
            return True
        if _ask(self, "Bu görevlendirme henüz onaylanmadı. Onaylayıp istatistiğe işleyelim mi?\n\n"
                      "Hayır derseniz çıktı alınır ama istatistiğe işlenmez."):
            self.approve()
        return True

    def _store(self):
        """Günün görevlendirmesini istatistiğe kaydeder (aynı gün tekrar kaydedilirse yenilenir)."""
        self.db.save_day(self.date_iso(), self.plan.absences, self.plan.assignments)
        self.win.statusBar().showMessage(
            "✅ Görevlendirme istatistiğe kaydedildi. Sonraki görevlendirmeler bu kayda göre "
            "adil dağıtılır.", 6000)
        self.win.write_excel()

    def delete(self):
        if _ask(self, "Bu günün devamsızlık ve görevlendirme kaydı silinsin mi?\n\nBu günün sayıları istatistiklerden de düşülür."):
            self.db.delete_day(self.date_iso())
            self.win.write_excel()
            self.refresh()

    def to_excel(self):
        if not self.plan:
            return
        path, _ = QFileDialog.getSaveFileName(
            self, "Excel olarak kaydet", f"gorevlendirme_{self.date_iso()}.xlsx", "Excel (*.xlsx)")
        if path:
            self._ensure_approved()
            export.save_day_xlsx(path, self.plan, self.win.school_header())
            QDesktopServices.openUrl(QUrl.fromLocalFile(path))

    def print_(self):
        if not self.plan:
            _info(self, "Önce gelmeyen öğretmenleri seçip 'Görevlendirmeyi Oluştur'a basın.")
            return
        self._ensure_approved()
        # Öğretmen programı satırları sayfa arasında bölünmesin
        doc = fit_day_document(self.plan, self.win.school_info(), os.path.exists(logo_path()),
                               self.schedule)
        preview(self, doc, f"gorevlendirme_{self.date_iso()}.pdf")


class ScheduleTab(QWidget):
    def __init__(self, win):
        super().__init__()
        self.win, self.db = win, win.db
        self.version = None
        self.data = {}

        self.versions = QTableWidget(0, 4)
        self.versions.setHorizontalHeaderLabels(["Geçerlilik", "Dosya", "Öğretmen", "Yüklenme"])
        align(self.versions, 0, 2, 3)
        self.versions.verticalHeader().hide()
        _stretch(self.versions, 1)
        self.versions.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.versions.setSelectionMode(QAbstractItemView.SingleSelection)
        self.versions.itemSelectionChanged.connect(self._select_version)
        load = variant(QPushButton("📄  Yeni ders programı yükle (PDF / Excel)"), "primary")
        load.clicked.connect(self.load_file)
        rem = variant(QPushButton("Seçili sürümü sil"), "danger")
        rem.clicked.connect(self.delete_version)

        self.teachers = QListWidget()
        self.teachers.setObjectName("list")
        self.teachers.currentTextChanged.connect(self._select_teacher)
        add_t = QPushButton("Öğretmen ekle")
        add_t.clicked.connect(self.add_teacher)
        ren_t = QPushButton("Yeniden adlandır")
        ren_t.clicked.connect(self.rename_teacher)
        del_t = QPushButton("Sil")
        del_t.clicked.connect(self.delete_teacher)

        self.grid = WeekGrid(8, editable=True)
        save = QPushButton("Bu öğretmenin programını kaydet")
        save.clicked.connect(self.save_teacher)

        # Kademe 1: yüklü programlar
        load.setMinimumHeight(46)
        load.setStyleSheet("font-size:11pt")
        view = QPushButton("👁  Programı görüntüle / düzenle")
        view.setMinimumHeight(46)
        view.clicked.connect(self.open_view)
        rem.setMinimumHeight(46)
        self.versions.setMinimumHeight(130)
        card1 = QGroupBox("Yüklü ders programları")
        c1 = QVBoxLayout(card1)
        hint = QLabel("Program değiştiğinde yeni dosyayı yükleyip <b>geçerlilik tarihini</b> girin. "
                      "Önceki günler eski programla, sonraki günler yeni programla hesaplanır.")
        hint.setWordWrap(True)
        c1.addWidget(hint)
        bh = QHBoxLayout()
        bh.addWidget(load, 1)
        bh.addWidget(view, 1)
        bh.addWidget(rem)
        c1.addLayout(bh)
        c1.addWidget(self.versions, 1)
        self.miss_main = QLabel()
        self.miss_main.setWordWrap(True)
        self.miss_main.setObjectName("warnBox")
        self.miss_main.hide()
        c1.addWidget(self.miss_main)

        # İdareciler: ders programı yüklendikten hemen sonra belirlenir
        self.admins_label = QLabel()
        self.admins_label.setWordWrap(True)
        edit_admins = variant(QPushButton("👤  İdarecileri belirle…"), "primary")
        edit_admins.clicked.connect(self.edit_admins)
        card2 = QGroupBox("İdareciler — okul müdürü ve müdür yardımcıları (otomatik görev verilmez)")
        c2 = QHBoxLayout(card2)
        c2.addWidget(self.admins_label, 1)
        c2.addWidget(edit_admins, 0, Qt.AlignVCenter)

        p1 = QWidget()
        v1 = QVBoxLayout(p1)
        v1.setContentsMargins(0, 0, 0, 0)
        v1.addWidget(card1, 1)
        v1.addWidget(card2)

        # Kademe 2: öğretmenler ve haftalık program (tam sayfa)
        back = QPushButton("←  Geri")
        back.clicked.connect(lambda: self.steps.setCurrentIndex(0))
        self.view_title = QLabel()
        self.view_title.setStyleSheet("font-size:12pt; font-weight:700")
        th = QHBoxLayout()
        th.addWidget(back)
        th.addWidget(self.view_title, 1)
        self.teachers.setFixedWidth(260)
        tl = QVBoxLayout()
        tl.addWidget(QLabel("<b>Öğretmenler</b>"))
        tl.addWidget(self.teachers, 1)
        for b in (add_t, ren_t, del_t):
            tl.addWidget(b)
        self.miss_info = QLabel()
        self.miss_info.setWordWrap(True)
        self.miss_info.setObjectName("warnBox")
        gr = QVBoxLayout()
        gr.addWidget(self.miss_info)
        gr.addWidget(QLabel("<b>Haftalık program</b> — hücreye çift tıklayıp düzenleyebilirsiniz"))
        gr.addWidget(self.grid, 1)
        gr.addWidget(save)
        body = QHBoxLayout()
        body.addLayout(tl)
        body.addLayout(gr, 1)
        p2 = QWidget()
        v2 = QVBoxLayout(p2)
        v2.setContentsMargins(0, 0, 0, 0)
        v2.addLayout(th)
        v2.addLayout(body, 1)

        self.steps = QStackedWidget()
        self.steps.addWidget(p1)
        self.steps.addWidget(p2)
        v = QVBoxLayout(self)
        v.setContentsMargins(0, 0, 0, 0)
        v.addWidget(page_header("📚  Ders Programı",
                                "Okulun ders programını PDF veya Excel olarak yükleyin; okunanları "
                                "kontrol edip gerekirse düzeltin."))
        v.addWidget(self.steps, 1)
        self.refresh()

    def refresh_admins(self):
        admins = sorted(self.db.admins(), key=normalize)
        if not self.db.all_teachers():
            self.admins_label.setText("Ders programını yükledikten sonra idarecileri (okul müdürü, "
                                      "müdür yardımcıları) burada işaretleyin.")
        elif admins:
            self.admins_label.setText(admin_roles_text(self.db) + "<br>Bu kişilere otomatik "
                                      "ders görevlendirmesi ve otomatik nöbet verilmez; "
                                      "görevlendirme ekranında elle seçilebilirler.")
        else:
            self.admins_label.setText("<span style='color:#dc2626'><b>Henüz idareci "
                                      "belirlenmedi.</b> Belirlemezseniz idarecilere de ders "
                                      "görevlendirilir.</span>")

    def edit_admins(self):
        if not self.db.all_teachers():
            _info(self, "Önce ders programını yükleyin.")
            return
        if run_dialog(AdminsDialog(self, self.db)):
            self.refresh_admins()
            self.win.school_tab.refresh()
            self.win.changed()

    def open_view(self):
        if not self.version:
            self.refresh()
        if not self.version:
            _info(self, "Önce ders programını yükleyin.")
            return
        r = self.versions.currentRow()
        if r >= 0:
            self.view_title.setText(f"Ders programı — {self.versions.item(r, 0).text()} "
                                    f"tarihinden itibaren ({self.versions.item(r, 1).text()})")
        self.steps.setCurrentIndex(1)

    def refresh(self, select=None):
        if hasattr(self, "admins_label"):
            self.refresh_admins()
        rows = self.db.schedule_versions()
        self.versions.blockSignals(True)
        self.versions.setRowCount(len(rows))
        for r, (vid, vf, src, imp, n) in enumerate(rows):
            self.versions.setItem(r, 0, _item(_qdate(vf).toString("dd.MM.yyyy")))
            self.versions.item(r, 0).setData(Qt.UserRole, vid)
            self.versions.setItem(r, 1, _item(src))
            self.versions.setItem(r, 2, _item(n))
            self.versions.setItem(r, 3, _item((imp or "").replace("T", " ")))
        self.versions.blockSignals(False)
        if rows:
            idx = next((r for r, row in enumerate(rows) if row[0] == select), 0)
            self.versions.selectRow(idx)
            self._select_version()
        else:
            self.version, self.data = None, {}
            self.teachers.clear()
            self.grid.load({})

    def _select_version(self):
        r = self.versions.currentRow()
        if r < 0:
            return
        self.version = self.versions.item(r, 0).data(Qt.UserRole)
        self.data = self.db.schedule(self.version)
        cur = self.teachers.currentItem().text() if self.teachers.currentItem() else None
        self.teachers.clear()
        self.teachers.addItems(sorted(self.data, key=normalize))
        self._mark_missing()
        text = missing_class_text(self.data)
        self.miss_main.setText(text.replace("Sarı hücrelere", "<b>Programı görüntüle / "
                                            "düzenle</b> ekranında sarı hücrelere"))
        self.miss_main.setVisible(bool(text))
        items = self.teachers.findItems(cur or "", Qt.MatchExactly) if cur else []
        self.teachers.setCurrentRow(self.teachers.row(items[0]) if items else 0)

    def _mark_missing(self):
        miss = missing_class(self.data)
        for i in range(self.teachers.count()):
            it = self.teachers.item(i)
            if it.text() in miss:
                it.setBackground(YELLOW)
                it.setToolTip(f"{len(miss[it.text()])} dersinde sınıf yazmıyor")
            else:
                it.setBackground(QColor(0, 0, 0, 0))
                it.setToolTip("")
        text = missing_class_text(self.data)
        self.miss_info.setText(text)
        self.miss_info.setVisible(bool(text))

    def _select_teacher(self, name):
        self.grid.load(self.data.get(name, {}))

    def load_file(self):
        path, _ = QFileDialog.getOpenFileName(self, "Ders programı seç", "", FILE_FILTER)
        if not path:
            return
        QApplication.setOverrideCursor(Qt.WaitCursor)
        try:
            res = parse_schedule(read_document(path), self.db.all_teachers())
        except ReadError as e:
            _warn(self, "Okunamadı", str(e))
            return
        finally:
            QApplication.restoreOverrideCursor()
        if not res.teachers:
            _warn(self, "Program bulunamadı", "\n".join(res.warnings))
            return
        dlg = ScheduleReviewDialog(self, res, path, dt.date.today().isoformat())
        if run_dialog(dlg):
            vid = self.db.save_schedule(dlg.teachers(), dlg.valid_from(), os.path.basename(path))
            self.refresh(select=vid)
            self.win.changed()
            self.win.next_step("Ders programı kaydedildi.")

    def delete_version(self):
        if self.version and _ask(self, "Seçili program sürümü silinsin mi?"):
            self.db.delete_schedule_version(self.version)
            self.refresh()
            self.win.changed()

    def _store(self):
        self.db.replace_schedule_entries(self.version, self.data)
        self.win.changed()

    def save_teacher(self):
        it = self.teachers.currentItem()
        if not it or not self.version:
            return
        self.data[it.text()] = self.grid.slots()
        self._store()
        self._mark_missing()
        self.win.statusBar().showMessage(f"{it.text()} programı kaydedildi.", 3000)

    def _ask_name(self, title, text=""):
        dlg = QDialog(self)
        dlg.setWindowTitle(title)
        edit = QLineEdit(text)
        bb = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        bb.button(QDialogButtonBox.Ok).setText("Tamam")
        bb.button(QDialogButtonBox.Cancel).setText("Vazgeç")
        bb.accepted.connect(dlg.accept)
        bb.rejected.connect(dlg.reject)
        v = QVBoxLayout(dlg)
        v.addWidget(QLabel("Ad Soyad:"))
        v.addWidget(edit)
        v.addWidget(bb)
        return clean_name(edit.text()) if run_dialog(dlg) else None

    def add_teacher(self):
        if not self.version:
            self.version = self.db.save_schedule({}, dt.date.today().isoformat(), "elle")
        name = self._ask_name("Öğretmen ekle")
        if name and name not in self.data:
            self.data[name] = {}
            self._store()
            self.refresh(select=self.version)
            self.teachers.setCurrentItem(self.teachers.findItems(name, Qt.MatchExactly)[0])

    def rename_teacher(self):
        it = self.teachers.currentItem()
        if not it:
            return
        old = it.text()
        name = self._ask_name("Yeniden adlandır", old)
        if name and name != old:
            merged = self.data.pop(old)
            self.data.setdefault(name, {}).update(merged)
            self._store()
            self.refresh(select=self.version)

    def delete_teacher(self):
        it = self.teachers.currentItem()
        if it and _ask(self, f"{it.text()} bu program sürümünden silinsin mi?"):
            self.data.pop(it.text(), None)
            self._store()
            self.refresh(select=self.version)


class DutyTab(QWidget):
    def __init__(self, win):
        super().__init__()
        self.win, self.db = win, win.db
        self.months = QTableWidget(0, 2)
        self.months.setHorizontalHeaderLabels(["Ay", "Öğretmen"])
        align(self.months, 1)
        self.months.verticalHeader().hide()
        _stretch(self.months, 0)
        self.months.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.months.setSelectionMode(QAbstractItemView.SingleSelection)
        self.months.itemSelectionChanged.connect(self._show)
        # Üç yol: otomatik oluştur / hazır çizelgeyi yükle / elle gir
        auto = variant(QPushButton("⚙  Otomatik oluştur\n"
                                   "     Program adil dağıtsın"), "step")
        auto.clicked.connect(lambda: self.open_nobetmatik())
        load = variant(QPushButton("📄  Hazır çizelgem var\n"
                                   "     PDF / Excel yükle"), "step")
        load.clicked.connect(self.load_file)
        manual = variant(QPushButton("✏  Elle gir / düzenle\n"
                                     "     Kendiniz yazın"), "step")
        manual.clicked.connect(self.edit)
        prt = QPushButton("🖨  Yazdır")
        prt.clicked.connect(self.print_month)
        rem = variant(QPushButton("Seçili ayı sil"), "danger")
        rem.clicked.connect(self.delete)

        self.table = QTableWidget(0, 6)
        self.table.setHorizontalHeaderLabels(["Hafta"] + DAYS)
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.table.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeToContents)
        self.table.setWordWrap(True)
        self.dated = QLabel()
        self.dated.setWordWrap(True)

        # Kademe 1: yüklü aylar
        view = QPushButton("👁  Görüntüle")
        view.clicked.connect(self.open_view)
        for b in (view, prt, rem):
            b.setMinimumHeight(40)
        self.months.setMinimumHeight(110)
        choose = QGroupBox("Bu ayın nöbet çizelgesini nasıl hazırlamak istersiniz?")
        ch = QHBoxLayout(choose)
        for b in (auto, load, manual):
            b.setMinimumHeight(64)
            ch.addWidget(b, 1)
        card1 = QGroupBox("Yüklü nöbet çizelgeleri")
        c1 = QVBoxLayout(card1)
        bh = QHBoxLayout()
        bh.addWidget(view, 1)
        bh.addWidget(prt, 1)
        bh.addWidget(rem)
        c1.addLayout(bh)
        c1.addWidget(self.months, 1)
        hint = QLabel("Görevlendirme, seçilen tarihin ayına ait çizelgeyi kullanır.")
        hint.setObjectName("subtitle")
        c1.addWidget(hint)
        p1 = QWidget()
        v1 = QVBoxLayout(p1)
        v1.setContentsMargins(0, 0, 0, 0)
        v1.addWidget(choose)
        v1.addWidget(card1, 1)

        # Kademe 2: haftalık nöbetçiler (tam sayfa)
        back = QPushButton("←  Geri")
        back.clicked.connect(lambda: self.steps.setCurrentIndex(0))
        self.view_title = QLabel()
        self.view_title.setStyleSheet("font-size:12pt; font-weight:700")
        th = QHBoxLayout()
        th.addWidget(back)
        th.addWidget(self.view_title, 1)
        self.table.setMinimumHeight(220)
        p2 = QWidget()
        v2 = QVBoxLayout(p2)
        v2.setContentsMargins(0, 0, 0, 0)
        v2.addLayout(th)
        v2.addWidget(self.table, 1)
        v2.addWidget(self.dated)

        # Nöbetmatik sayfası açılınca kurulur, kapanınca silinir: yüzlerce denetim sürekli
        # bellekte durup programı (ve tema değişimini) yavaşlatmasın.
        self.matik = None
        self.steps = QStackedWidget()
        self.steps.addWidget(p1)
        self.steps.addWidget(p2)
        v = QVBoxLayout(self)
        v.setContentsMargins(0, 0, 0, 0)
        v.addWidget(page_header("🗓  Nöbet Çizelgesi",
                                "Ders programındaki öğretmenlerden her ay adil bir nöbet çizelgesi "
                                "oluşturun; görevlendirme bu çizelgeyi kullanır."))
        v.addWidget(self.steps, 1)
        self.refresh()

    def open_nobetmatik(self, month=None):
        if not self.db.schedule_versions():
            _info(self, "Önce Ders Programı sayfasından ders programını yükleyin; otomatik dağıtım öğretmenleri ve "
                        "derslerini oradan alır.")
            return
        self.win.tabs.setCurrentWidget(self)
        if self.matik is None:
            from .ui_nobetmatik import NobetmatikPage
            self.matik = NobetmatikPage(self.win, self._matik_done, self._close_matik)
            self.steps.addWidget(self.matik)
        self.matik.start(month or dt.date.today().strftime("%Y-%m"))
        self.steps.setCurrentWidget(self.matik)

    def _close_matik(self):
        self.steps.setCurrentIndex(0)
        if self.matik is not None:
            self.steps.removeWidget(self.matik)
            self.matik.deleteLater()
            self.matik = None

    def _matik_done(self, month):
        self._close_matik()
        self.refresh(select=month)
        self.win.changed()
        self.win.next_step(f"{export.month_tr(month)} nöbet çizelgesi kaydedildi.")

    def _month_entries(self, m):
        """Ayın nöbetlerini tarihli olarak verir (gün bazlı çizelgeler her haftaya yayılır)."""
        from .nobetmatik import month_weeks
        out = []
        dates = [d for _mon, ds in month_weeks(m) for d in ds]
        for wd, date, t, p in self.db.duty_rows(m):
            if date:
                out.append(DutyEntry(wd, t, p, date))
            else:
                out += [DutyEntry(wd, t, p, d.isoformat()) for d in dates if d.weekday() == wd]
        return out

    def print_month(self):
        m = self.month()
        if not m:
            _info(self, "Yazdırılacak çizelge yok.")
            return
        from .ui_nobetmatik import print_duty
        cfg = self.db.nobetmatik_config()
        exempt = cfg.get("exempt") if cfg.get("month") == m else None
        places = [p for p, _n in cfg.get("places", [])] if cfg.get("month") == m else None
        print_duty(self, self.win, m, self._month_entries(m), exempt, places)

    def open_view(self):
        if not self.month():
            self.refresh()
        m = self.month()
        if not m:
            _info(self, "Önce nöbet çizelgesini yükleyin.")
            return
        self.view_title.setText(f"{export.month_tr(m)} nöbet çizelgesi")
        self._show()
        self.steps.setCurrentIndex(1)

    def month(self):
        r = self.months.currentRow()
        return self.months.item(r, 0).data(Qt.UserRole) if r >= 0 else None

    def refresh(self, select=None):
        rows = self.db.duty_months()
        self.months.blockSignals(True)
        self.months.setRowCount(len(rows))
        for r, (m, n) in enumerate(rows):
            self.months.setItem(r, 0, _item(export.month_tr(m)))
            self.months.item(r, 0).setData(Qt.UserRole, m)
            self.months.setItem(r, 1, _item(n))
        self.months.blockSignals(False)
        if rows:
            self.months.selectRow(next((r for r, row in enumerate(rows) if row[0] == select), 0))
        self._show()

    def _show(self):
        m = self.month()
        rows = self.db.duty_rows(m) if m else []
        cfg = self.db.nobetmatik_config()
        places = [p for p, _n in cfg.get("places", [])] if cfg.get("month") == m else []
        places = [p for p in places if any(r[3] == p for r in rows)]
        for _wd, _d, _t, p in rows:
            if p and p not in places and not is_admin_place(p):
                places.append(p)
        has_admin = any(is_admin_place(p or "") for *_x, p in rows)
        cols = places + (["Nöbetçi idareci"] if has_admin else [])
        dated = any(d for _wd, d, _t, _p in rows)
        grid_rows = []
        if dated:
            by_date = {}
            for wd, d, t, p in rows:
                if d:
                    by_date.setdefault(d, []).append((t, p))
            for d in sorted(by_date):
                q = _qdate(d)
                label = f"{DAYS[q.dayOfWeek() - 1]}\n{q.toString('dd.MM.yyyy')}"
                cells = [[t for t, p in by_date[d] if p == col] for col in places]
                if has_admin:
                    cells.append([t for t, p in by_date[d] if is_admin_place(p or "")])
                grid_rows.append((label, cells))
        else:
            for wd in range(5):
                day_rows = [(t, p) for w, _d, t, p in rows if w == wd]
                cells = [[t for t, p in day_rows if p == col] for col in places]
                if has_admin:
                    cells.append([t for t, p in day_rows if is_admin_place(p or "")])
                if not places and not has_admin:
                    cells = [[t for t, _p in day_rows]]
                grid_rows.append((f"{DAYS[wd]}\n(her hafta)", cells))
        fill_duty_grid(self.table, ["Gün"] + (cols or ["Nöbetçiler"]), grid_rows)
        self.dated.setText("İdareciler otomatik görevlendirilmez; görevlendirme ekranında elle "
                           "seçilebilir."
                           if any(is_admin_place(r[3]) for r in rows) else "")

    def _teachers(self):
        return self.db.all_teachers()

    def load_file(self):
        teachers = self._teachers()
        if not teachers:
            _info(self, "Önce ders programını yükleyin. Nöbet çizelgesindeki isimler "
                        "ders programındaki öğretmenlerle eşleştirilir.")
            return
        path, _ = QFileDialog.getOpenFileName(self, "Nöbet çizelgesi seç", "", FILE_FILTER)
        if not path:
            return
        QApplication.setOverrideCursor(Qt.WaitCursor)
        try:
            res = parse_duty(read_document(path), teachers)
        except ReadError as e:
            _warn(self, "Okunamadı", str(e))
            return
        finally:
            QApplication.restoreOverrideCursor()
        month = _guess_month(res.entries)
        dlg = DutyEditDialog(self, res.entries, teachers, month, res.unmatched, res.warnings)
        if run_dialog(dlg):
            self._save(dlg)

    def edit(self):
        """Elle giriş: seçili ayın çizelgesini (yoksa boş bir çizelgeyi) düzenler."""
        teachers = self._teachers()
        if not teachers:
            _info(self, "Önce Ders Programı sayfasından ders programını yükleyin.")
            return
        m = self.month() or dt.date.today().strftime("%Y-%m")
        entries = [DutyEntry(wd, t, p, d) for wd, d, t, p in self.db.duty_rows(m)]
        if not entries:
            entries = [DutyEntry(0, "", "")]
        dlg = DutyEditDialog(self, entries, teachers, m)
        dlg.setWindowTitle("Nöbet çizelgesi — elle giriş / düzenleme")
        if run_dialog(dlg):
            self._save(dlg)

    def _save(self, dlg):
        m = dlg.month_key()
        if m in [r[0] for r in self.db.duty_months()] and m != self.month() and \
                not _ask(self, f"{export.month_tr(m)} için kayıtlı çizelge var. Üzerine yazılsın mı?"):
            return
        self.db.save_duty(m, dlg.entries())
        self.refresh(select=m)
        self.win.changed()
        self.win.next_step("Nöbet çizelgesi kaydedildi.")

    def delete(self):
        m = self.month()
        if m and _ask(self, f"{export.month_tr(m)} nöbet çizelgesi silinsin mi?"):
            self.db.delete_duty_month(m)
            self.refresh()
            self.win.changed()


class RecordsTab(QWidget):
    """Onaylanmış görevlendirmeler: solda günler, sağda seçilen günün ayrıntısı."""

    def __init__(self, win):
        super().__init__()
        self.win, self.db = win, win.db
        self.month = MonthPicker()
        self.month.changed.connect(self.refresh)
        self.all_year = QCheckBox("Bütün eğitim yılı")
        self.all_year.toggled.connect(self.refresh)

        # Sol: günler (sade)
        self.table = QTableWidget(0, 3)
        self.table.setHorizontalHeaderLabels(["Tarih", "Gelmeyen öğretmen", "Durum"])
        _stretch(self.table, 1)
        align(self.table, 0)
        for c, w in ((0, 100), (2, 84)):
            self.table.horizontalHeader().setSectionResizeMode(c, QHeaderView.Fixed)
            self.table.setColumnWidth(c, w)
        self.table.verticalHeader().hide()
        self.table.setShowGrid(False)
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.table.setSelectionMode(QAbstractItemView.SingleSelection)
        self.table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.table.setWordWrap(True)
        self.table.itemSelectionChanged.connect(self._show_detail)
        self.table.doubleClicked.connect(lambda _i: self.open_day())
        left = QGroupBox("Gelmeyen öğretmenler")
        ll = QVBoxLayout(left)
        ll.setContentsMargins(12, 18, 12, 12)
        ll.addWidget(self.table)

        # Sağ: seçilen günün dersleri
        self.detail = QTableWidget(0, 4)
        self.detail.setHorizontalHeaderLabels(["Saat", "Sınıf", "Ders", "Derse giren"])
        _stretch(self.detail, 3)
        align(self.detail, 0, 1)
        for c, w in ((0, 64), (1, 56), (2, 110)):
            self.detail.horizontalHeader().setSectionResizeMode(c, QHeaderView.Fixed)
            self.detail.setColumnWidth(c, w)
        self.detail.verticalHeader().hide()
        self.detail.setShowGrid(False)
        self.detail.setSelectionMode(QAbstractItemView.NoSelection)
        self.detail.setFocusPolicy(Qt.NoFocus)
        self.detail.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.detail_box = QGroupBox("Seçili gün")
        dl = QVBoxLayout(self.detail_box)
        dl.setContentsMargins(12, 18, 12, 12)
        dl.addWidget(self.detail)
        open_b = variant(QPushButton("📂  Aç"), "primary")
        open_b.clicked.connect(self.open_day)
        prt = QPushButton("🖨  Yazdır")
        prt.clicked.connect(lambda: self.open_day(print_it=True))
        rem = variant(QPushButton("🗑  Kaydı sil"), "danger")
        rem.setToolTip("Yalnızca seçili öğretmenin bu günkü kaydını siler")
        rem.clicked.connect(self.delete_day)
        db_ = QHBoxLayout()
        db_.addWidget(open_b)
        db_.addWidget(prt)
        db_.addStretch()
        db_.addWidget(rem)
        dl.addLayout(db_)

        rem_all = variant(QPushButton("Listedeki bütün kayıtları sil…"), "danger")
        rem_all.clicked.connect(self.delete_all)
        self.info = QLabel()
        self.info.setObjectName("subtitle")
        self.info.setWordWrap(True)
        top = QHBoxLayout()
        top.addWidget(QLabel("<b>Ay:</b>"))
        top.addWidget(self.month)
        top.addSpacing(12)
        top.addWidget(self.all_year)
        top.addStretch()
        top.addWidget(rem_all)
        body = QHBoxLayout()
        body.setSpacing(14)
        left.setMinimumWidth(340)
        body.addWidget(left, 5)
        body.addWidget(self.detail_box, 6)
        v = QVBoxLayout(self)
        v.setContentsMargins(0, 0, 0, 0)
        v.addWidget(page_header("📋  Görevlendirme Kayıtları",
                                "Onaylanan her günün görevlendirmesi burada saklanır. Bir günü "
                                "silerseniz istatistiklerden de silinir."))
        v.addLayout(top)
        v.addWidget(self.info)
        v.addLayout(body, 1)
        keep_still(self.table, self.detail)
        self.dates = []
        self.rows = []

    def _range(self):
        if self.all_year.isChecked():
            start = self.db.school_year_start(self.month.key() + "-01")
            return start, f"{int(start[:4]) + 1}-08-31"
        m = self.month.key()
        return m + "-01", m + "-31"

    def refresh(self):
        days = self.db.saved_days(*self._range())
        # Her gelmeyen öğretmen ayrı satır: (tarih, öğretmen, durum, dersleri)
        self.rows = []
        for date, absent, rows in days:
            people = sorted(set(absent) | {x[1] for x in rows}, key=normalize)
            for t in people:
                self.rows.append((date, t, absent.get(t, ""), [x for x in rows if x[1] == t]))
        self.dates = [r[0] for r in self.rows]
        self.table.setRowCount(len(self.rows))
        lessons = 0
        for r, (date, t, why, rows) in enumerate(self.rows):
            d = dt.date.fromisoformat(date)
            lessons += len(rows)
            day = DAYS[d.weekday()] if d.weekday() < 5 else ""
            for c, val in enumerate((f"{d.strftime('%d.%m.%Y')}\n{day}", t, why)):
                self.table.setItem(r, c, _item(val))
        self.table.resizeRowsToContents()
        for r in range(self.table.rowCount()):
            self.table.setRowHeight(r, self.table.rowHeight(r) + 10)
        period = (f"{self._range()[0][:4]}–{int(self._range()[0][:4]) + 1} eğitim yılı"
                  if self.all_year.isChecked() else export.month_tr(self.month.key()))
        self.info.setText(f"{period}: <b>{len(days)}</b> gün, <b>{len(self.rows)}</b> gelmeyen "
                          f"öğretmen, <b>{lessons}</b> boş ders görevlendirmesi kayıtlı. "
                          "Ayrıntısını görmek için soldan bir satır seçin."
                          if self.rows else f"{period} için onaylanmış görevlendirme yok.")
        if self.rows:
            self.table.selectRow(0)
        self._show_detail()

    def _show_detail(self):
        r = self.table.currentRow()
        self.detail.clearSpans()
        self.detail.setRowCount(0)
        if r < 0 or r >= len(self.rows):
            self.detail_box.setTitle("Seçili kayıt")
            self.detail_box.setEnabled(False)
            return
        self.detail_box.setEnabled(True)
        date, who, why, rows = self.rows[r]
        d = dt.date.fromisoformat(date)
        self.detail_box.setTitle(f"{who} · {why} — {d.strftime('%d.%m.%Y')} "
                                 f"{DAYS[d.weekday()]}, {len(rows)} ders")
        for hour, _who, info, assigned, kind, note in sorted(rows):
            sinif, ders = split_info(info or "")
            ders = lesson_label(ders)
            i = self.detail.rowCount()
            self.detail.insertRow(i)
            giren = _item(assigned or "— atanamadı —")
            if not assigned:
                giren.setBackground(RED)
            elif kind == "joker":
                giren.setBackground(PURPLE)
                giren.setText(f"{assigned}  (joker)")
            elif kind != "nobetci":
                giren.setBackground(YELLOW)
                giren.setText(f"{assigned}  (nöbetçi değil, o saatte boş)")
            for c, val in enumerate((f"{hour}. ders", sinif or "—", ders or "—")):
                self.detail.setItem(i, c, _item(val))
            self.detail.setItem(i, 3, giren)

    def _selected(self):
        r = self.table.currentRow()
        if r < 0 or r >= len(self.dates):
            _info(self, "Önce listeden bir gün seçin.")
            return None
        return self.dates[r]

    def open_day(self, print_it=False):
        date = self._selected()
        if not date:
            return
        daily = self.win.daily
        daily.date.setDate(QDate.fromString(date, "yyyy-MM-dd"))
        self.win.tabs.setCurrentWidget(daily)
        daily.refresh()
        if daily.plan:
            daily.go(2)
            if print_it:
                daily.print_()

    def delete_day(self):
        date = self._selected()
        if not date:
            return
        _d, who, why, rows = self.rows[self.table.currentRow()]
        if _ask(self, f"{dt.date.fromisoformat(date).strftime('%d.%m.%Y')} — {who} ({why}) "
                      f"kaydı ve {len(rows)} dersin görevlendirmesi silinsin mi?\n\nO günün "
                      "diğer öğretmenleri silinmez. Bu sayılar istatistiklerden de düşülür."):
            self.db.delete_absent(date, who)
            self.win.write_excel()
            self.refresh()

    def delete_all(self):
        if not self.dates:
            return
        if _ask(self, f"Listedeki {len(set(self.dates))} günün bütün görevlendirme kayıtları "
                      "silinsin mi?\n\nİstatistikler de sıfırlanır. Bu işlem geri alınamaz "
                      "(isterseniz önce Ayarlar'dan yedek alın).", "Kayıtları sil"):
            for d in set(self.dates):
                self.db.delete_day(d)
            self.win.write_excel()
            self.refresh()


class ReportTab(QWidget):
    """Adalet raporu: herkese gösterildiğinde dağıtımın adil olduğu açıkça görülsün."""

    HEADERS = ["Öğretmen", "Nöbetçi olduğu gün\n(çizelgeye göre)", "Derse girebileceği\nboş saat",
               "Nöbetçiyken\ngirdiği", "Nöbetçi değilken\ngirdiği", "Toplam", ""]

    def __init__(self, win):
        super().__init__()
        self.win, self.db = win, win.db
        self.period = QComboBox()
        self.period.addItems(["Seçilen ay", "Eğitim yılı (Eylül'den bu yana)"])
        self.period.setMinimumWidth(240)
        self.period.currentIndexChanged.connect(self.refresh)
        self.month = MonthPicker()
        self.month.changed.connect(self.refresh)
        prt = variant(QPushButton("🖨  Raporu yazdır"), "primary")
        prt.clicked.connect(self.print_)
        xlsx = QPushButton("Excel'e aktar")
        xlsx.clicked.connect(self.to_excel)
        top = QHBoxLayout()
        top.addWidget(QLabel("<b>Dönem:</b>"))
        top.addWidget(self.period)
        top.addWidget(self.month)
        top.addStretch()
        top.addWidget(prt)
        top.addWidget(xlsx)
        self.verdict = QLabel()
        self.verdict.setStyleSheet("font-size:14pt; font-weight:700")
        self.verdict_text = QLabel()
        self.verdict_text.setWordWrap(True)
        card = QFrame()
        card.setObjectName("card")
        cl = QVBoxLayout(card)
        cl.setContentsMargins(16, 10, 16, 10)
        cl.addWidget(self.verdict)
        cl.addWidget(self.verdict_text)
        self.table = QTableWidget(0, len(self.HEADERS))
        self.table.setHorizontalHeaderLabels(self.HEADERS)
        align(self.table, 1, 2, 3, 4, 5)
        h = self.table.horizontalHeader()
        h.setSectionResizeMode(0, QHeaderView.Stretch)
        for c in range(1, 6):
            h.setSectionResizeMode(c, QHeaderView.ResizeToContents)
        h.setSectionResizeMode(6, QHeaderView.Stretch)
        self.table.verticalHeader().hide()
        self.table.verticalHeader().setDefaultSectionSize(theme.ROW_HEIGHT)
        self.table.setSelectionMode(QAbstractItemView.NoSelection)
        self.table.setFocusPolicy(Qt.NoFocus)
        self.table.setShowGrid(False)
        self.table.setMinimumHeight(180)
        rule = QLabel("<b>Kural:</b> Gelmeyen öğretmenin her dersi için, o saatte <b>kendi dersi "
                      "olmayan</b> nöbetçiler arasından <b>bu yıl en az derse girmiş olan</b> seçilir; "
                      "boş nöbetçi yoksa o saatte dersi olmayan diğer öğretmenlerden seçilir. "
                      "<b>Derse girebileceği boş saat</b>: bir ders boşa düştüğünde öğretmenin o gün "
                      "nöbetçi olup kendi dersinin olmadığı durum sayısı.")
        rule.setWordWrap(True)
        rule.setObjectName("subtitle")
        v = QVBoxLayout(self)
        v.setContentsMargins(0, 0, 0, 0)
        v.addWidget(page_header("Aylık Rapor"))
        v.addLayout(top)
        v.addWidget(card)
        v.addWidget(self.table, 1)
        v.addWidget(rule)
        self.stats = {}

    def key(self):
        return self.month.key()

    def _range(self):
        m = self.key()
        y, mo = (int(x) for x in m.split("-"))
        import calendar
        end = f"{m}-{calendar.monthrange(y, mo)[1]:02d}"
        end = min(end, dt.date.today().isoformat()) if m <= dt.date.today().strftime("%Y-%m") else end
        if self.period.currentIndex() == 1:
            start = Database.school_year_start(end)
            title = (f"{_qdate(start).toString('dd.MM.yyyy')} – "
                     f"{_qdate(end).toString('dd.MM.yyyy')} (eğitim yılı)")
        else:
            start = f"{m}-01"
            title = export.month_tr(m)
        return start, end, title

    def refresh(self):
        from PySide6.QtWidgets import QProgressBar
        self.month.setEnabled(True)
        start, end, title = self._range()
        self.title = title
        self.stats = self.db.fairness_stats(start, end)
        rows = sorted(((t, v) for t, v in self.stats.items()
                       if t != "__meta__" and (v["days"] or v["assigned"])),
                      key=lambda x: (-x[1]["assigned"], normalize(x[0])))
        if not any(v["assigned"] for _t, v in rows):
            rows = []   # görevlendirme yoksa (ör. silindiyse) tablo boş kalır
        self.rows = rows
        verdict, color, lines = export.fairness_summary(self.stats)
        if not rows:
            verdict, lines = "Bu dönemde kayıtlı görevlendirme yok", [
                "Görevlendirme yaptıkça bu tablo dolar. Silinen görevlendirmeler sayılmaz."]
        self.verdict.setText(f"{title}:  {verdict}")
        self.verdict.setStyleSheet(f"font-size:14pt; font-weight:700; color:{color}")
        total = sum(v["assigned"] for _t, v in rows)
        self.verdict_text.setText(f"Bu dönemde gelmeyen öğretmenlerin <b>{total}</b> dersi "
                                  "dolduruldu. " + " ".join(lines))
        exc = self.stats.get("__meta__", {}).get("exceptions", [])
        self.verdict_text.setToolTip("\n".join(f"{_qdate(d).toString('dd.MM')} {h}. ders: {t} "
                                               f"({why})" for d, h, t, why in exc))
        top = max((v["assigned"] for _t, v in rows), default=0) or 1
        self.table.setRowCount(len(rows))
        for i, (t, v) in enumerate(rows):
            self.table.setItem(i, 0, _item(t))
            for c, val in enumerate((v["days"], v["chances"], v["as_duty"], v["as_other"],
                                     v["assigned"]), start=1):
                it = _item(val)
                it.setTextAlignment(Qt.AlignCenter)
                if c == 5:
                    f = it.font()
                    f.setBold(True)
                    it.setFont(f)
                if c == 4 and val:
                    it.setBackground(YELLOW)
                self.table.setItem(i, c, it)
            bar = QProgressBar()
            bar.setRange(0, top)
            bar.setValue(v["assigned"])
            bar.setTextVisible(False)
            bar.setFixedHeight(12)
            bar.setStyleSheet(f"QProgressBar {{ border:none; background:transparent; }}"
                              f"QProgressBar::chunk {{ background:{theme.ACCENT}; border-radius:6px; }}")
            holder = QWidget()
            hl = QVBoxLayout(holder)
            hl.setContentsMargins(8, 0, 8, 0)
            hl.addWidget(bar)
            self.table.setCellWidget(i, 6, holder)

    def print_(self):
        self.refresh()
        doc = html_document(export.report_html(self.title, self.stats, self.win.school_info(),
                                               os.path.exists(logo_path())))
        preview(self, doc, "rapor.pdf")

    def to_excel(self):
        import openpyxl
        from openpyxl.styles import Font, PatternFill
        path, _ = QFileDialog.getSaveFileName(self, "Excel olarak kaydet",
                                              f"adalet_raporu_{self.key()}.xlsx", "Excel (*.xlsx)")
        if not path:
            return
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Rapor"
        ws.append([self.win.school_header()])
        ws.append([f"Boş ders görevlendirme raporu — {self.title}"])
        ws.append([export.fairness_summary(self.stats)[0]])
        ws.append([])
        ws.append([h.replace("\n", " ") for h in self.HEADERS[:6]])
        for c in ws[5]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="DDE4EE")
        for t, v in self.rows:
            ws.append([t, v["days"], v["chances"], v["as_duty"], v["as_other"], v["assigned"]])
        for col, w in zip("ABCDEF", (30, 14, 22, 18, 22, 10)):
            ws.column_dimensions[col].width = w
        wb.save(path)
        QDesktopServices.openUrl(QUrl.fromLocalFile(path))


ROLES = ["—", "Okul Müdürü", "Müdür Yardımcısı", "Diğer idareci"]


class AdminsDialog(QDialog):
    """İdarecileri ve görevlerini (okul müdürü / müdür yardımcısı) seçme.

    İdareci olarak seçilenlere otomatik ders görevlendirmesi yapılmaz."""

    def __init__(self, parent, db):
        super().__init__(parent)
        self.db = db
        self.setWindowTitle("İdareciler")
        self.resize(620, 620)
        teachers = db.all_teachers()
        selected = db.admins()
        matcher = NameMatcher(teachers)
        principal = matcher.match(db.get_setting("principal", "")) \
            if db.get_setting("principal", "").strip() else None
        vices = {matcher.match(n) for n in re.split(r"[,;\n]", db.get_setting("vice_principal", ""))
                 if n.strip()}
        self.search = QLineEdit()
        self.search.setPlaceholderText("🔍  Öğretmen ara…")
        chosen = set(selected) | vices | {principal}
        names = sorted(set(teachers) | set(selected),
                       key=lambda t: (t != principal, t not in vices, t not in chosen, normalize(t)))
        self.table = QTableWidget(len(names), 2)
        self.table.setHorizontalHeaderLabels(["Öğretmen", "Görevi"])
        align(self.table)
        self.table.horizontalHeader().setSectionResizeMode(0, QHeaderView.Stretch)
        self.table.horizontalHeader().setSectionResizeMode(1, QHeaderView.Fixed)
        self.table.setColumnWidth(1, 200)
        self.table.verticalHeader().hide()
        self.table.verticalHeader().setDefaultSectionSize(40)
        self.table.setSelectionMode(QAbstractItemView.NoSelection)
        self.table.setFocusPolicy(Qt.NoFocus)
        self.combos = []
        for r, t in enumerate(names):
            self.table.setItem(r, 0, _item(t))
            c = QComboBox()
            c.addItems(ROLES)
            if t == principal:
                c.setCurrentIndex(1)
            elif t in vices:
                c.setCurrentIndex(2)
            elif t in selected:
                c.setCurrentIndex(3)
            c.currentIndexChanged.connect(lambda _i, row=r: self._mark(row))
            self.table.setCellWidget(r, 1, c)
            self.combos.append((t, c))
            self._mark(r)
        self.search.textChanged.connect(self._filter)
        bb = QDialogButtonBox(QDialogButtonBox.Save | QDialogButtonBox.Cancel)
        bb.button(QDialogButtonBox.Save).setText("Kaydet")
        variant(bb.button(QDialogButtonBox.Save), "primary")
        bb.button(QDialogButtonBox.Cancel).setText("Vazgeç")
        bb.accepted.connect(self._save)
        bb.rejected.connect(self.reject)
        v = QVBoxLayout(self)
        v.addWidget(page_header("👤  İdareciler",
                                "Her idarecinin satırında <b>Görevi</b> kutusundan Okul Müdürü, "
                                "Müdür Yardımcısı veya Diğer idareci seçin. Seçilenlere otomatik "
                                "ders görevlendirmesi yapılmaz; isterseniz sonuç ekranında elle "
                                "seçersiniz."))
        v.addWidget(self.search)
        v.addWidget(self.table, 1)
        v.addWidget(bb)

    def _mark(self, row):
        on = self.combos[row][1].currentIndex() > 0
        it = self.table.item(row, 0)
        f = it.font()
        f.setBold(on)
        it.setFont(f)
        it.setBackground(theme.BLUE if on else QColor(0, 0, 0, 0))

    def _filter(self, text):
        key = normalize(text)
        for r, (t, _c) in enumerate(self.combos):
            self.table.setRowHidden(r, bool(key) and key not in normalize(t))

    def roles(self):
        return {t: ROLES[c.currentIndex()] for t, c in self.combos if c.currentIndex() > 0}

    def selected(self):
        return set(self.roles())

    def _save(self):
        roles = self.roles()
        principals = [t for t, r in roles.items() if r == "Okul Müdürü"]
        if len(principals) > 1:
            _info(self, "Yalnızca bir kişiyi Okul Müdürü seçebilirsiniz.")
            return
        # Önce okul bilgilerindeki müdür / müdür yardımcısı güncellenir, sonra idareci listesi
        self.db.set_setting("principal", principals[0] if principals else "")
        self.db.set_setting("vice_principal", ", ".join(
            t for t, r in roles.items() if r == "Müdür Yardımcısı"))
        self.db.set_admins(set(roles))
        self.db.set_setting("admins_confirmed", "1")
        self.accept()


class SettingsTab(QWidget):
    def __init__(self, win):
        super().__init__()
        self.win, self.db = win, win.db

        def card(title):
            g = QGroupBox(title)
            lay = QVBoxLayout(g)
            lay.setSpacing(8)
            return g, lay

        def row(*widgets, stretch_first=True):
            h = QHBoxLayout()
            for i, w in enumerate(widgets):
                h.addWidget(w, 1 if (i == 0 and stretch_first) else 0)
            if not stretch_first:
                h.addStretch()
            return h

        # Görünüm: renk kutucukları ve seçenek düğmeleri (açılır liste yerine tıkla-seç)
        gt, lt = card("Görünüm (tasarım)")
        self.accent_name = self.db.get_setting("theme_accent", "Mor")
        self.swatches = {}
        sw = QHBoxLayout()
        sw.setSpacing(8)
        for name, (color, *_rest) in theme.ACCENTS.items():
            b = QPushButton()
            b.setFixedSize(38, 38)
            b.setToolTip(name)
            b.setCursor(Qt.PointingHandCursor)
            b.clicked.connect(lambda _c=False, n=name: self._pick_accent(n))
            self.swatches[name] = (b, color)
            sw.addWidget(b)
        sw.addStretch()
        self._paint_swatches()

        def choices(key, options, default):
            box = QHBoxLayout()
            box.setSpacing(6)
            group = []
            current = self.db.get_setting(key, default)
            for opt in options:
                b = QPushButton(opt)
                b.setCheckable(True)
                b.setChecked(opt == current)
                b.setCursor(Qt.PointingHandCursor)
                b.setMinimumWidth(92)
                group.append(b)
                box.addWidget(b)
            for b in group:
                b.clicked.connect(lambda _c=False, btn=b, grp=group, k=key: self._choose(k, btn, grp))
            box.addStretch()
            self._style_choices(group)
            return box

        self.family = QComboBox()
        self.family.addItems(theme.FONTS)
        self.family.setCurrentText(self.db.get_setting("theme_family", "Segoe UI"))
        self.family.setFixedWidth(220)
        self.family.currentTextChanged.connect(
            lambda v: (self.db.set_setting("theme_family", v), self._theme_changed()))
        ft = QFormLayout()
        ft.setVerticalSpacing(10)
        ft.addRow("Tema rengi:", sw)
        ft.addRow("Görünüm:", choices("theme_mode", ["Açık", "Koyu"],
                                      "Koyu" if self.db.get_setting("theme_dark") == "1" else "Açık"))
        ft.addRow("Yazı tipi:", self.family)
        ft.addRow("Yazı boyutu:", choices("theme_font", list(theme.FONT_SIZES), "Normal"))
        ft.addRow("Satır aralığı:", choices("theme_density", list(theme.DENSITIES), "Normal"))
        ft.addRow("Köşeler:", choices("theme_radius", list(theme.RADII), "Normal"))
        ft.addRow("Arka plan:", choices("theme_bg", list(theme.BACKGROUNDS), "Sade"))
        lt.addLayout(ft)
        note = QLabel("Seçim anında uygulanır. Satır aralığı tablolara sayfa açılınca yansır.")
        note.setObjectName("subtitle")
        lt.addWidget(note)

        # Dağıtım
        self.cap = QSpinBox()
        self.cap.setRange(0, 8)
        self.cap.setSpecialValueText("Sınırsız")
        self.cap.setFixedWidth(120)
        self.cap.setValue(int(self.db.get_setting("max_per_day", "0") or 0))
        self.cap.valueChanged.connect(lambda v: self.db.set_setting("max_per_day", v))
        g1, l1 = card("Dağıtım kuralları")
        rules = QLabel("Önce o saatte dersi olmayan <b>nöbetçi öğretmenler</b>, bu eğitim yılında en "
                       "az derse girene öncelik verilerek görevlendirilir. Boş nöbetçi yoksa o saatte dersi "
                       "olmayan <b>diğer öğretmenler</b> görevlendirilir. İdareciler otomatik "
                       "görevlendirilmez; sonuç ekranında elle seçilebilir.")
        rules.setWordWrap(True)
        l1.addWidget(rules)
        l1.addLayout(row(QLabel("Bir nöbetçiye günde en fazla:"), self.cap, stretch_first=False))

        # Yeni ay
        self.month_alert = QCheckBox("Yeni ay başlayınca uyar ve yeni nöbet çizelgesi oluşturmayı öner")
        self.month_alert.setChecked(self.db.get_setting("month_alert", "1") == "1")
        self.month_alert.toggled.connect(
            lambda v: self.db.set_setting("month_alert", "1" if v else "0"))
        new_month = QPushButton("⚙  Yeni ayı şimdi başlat (otomatik nöbet çizelgesi)")
        new_month.clicked.connect(lambda: win.duty.open_nobetmatik())
        gm, lm = card("Yeni ay")
        mhint = QLabel("Program tarihi kontrol eder; yeni bir aya geçildiğinde o ayın nöbet "
                       "çizelgesini oluşturmanızı ister. İstatistikler eğitim-öğretim yılı boyunca "
                       "(Eylül–Ağustos) birikir, ay değişince sıfırlanmaz.")
        mhint.setWordWrap(True)
        lm.addWidget(mhint)
        lm.addWidget(self.month_alert)
        lm.addLayout(row(new_month, stretch_first=False))

        # Excel
        g4, l4 = card("Excel kayıt dosyası")
        hint = QLabel("Her görevlendirmeden sonra tüm kayıtlar bu Excel dosyasına yazılır:")
        hint.setWordWrap(True)
        l4.addWidget(hint)
        self.excel = QLabel(win.excel_path())
        self.excel.setTextInteractionFlags(Qt.TextSelectableByMouse)
        self.excel.setWordWrap(True)
        self.excel.setStyleSheet("font-weight:600")
        l4.addWidget(self.excel)
        open_x = QPushButton("Excel dosyasını aç")
        open_x.clicked.connect(self.open_excel)
        move_x = QPushButton("Konumu değiştir…")
        move_x.clicked.connect(self.move_excel)
        l4.addLayout(row(open_x, move_x, stretch_first=False))

        # Güncelleme
        self.upd = QLabel(f"Yüklü sürüm: <b>{__version__}</b>")
        self.upd.setWordWrap(True)
        self.auto = QCheckBox("Program açılırken yeni sürüm var mı diye bak")
        self.auto.setChecked(self.db.get_setting("auto_update", "1") == "1")
        self.auto.toggled.connect(lambda v: self.db.set_setting("auto_update", "1" if v else "0"))
        check = variant(QPushButton("🔄  Güncellemeleri denetle"), "primary")
        check.clicked.connect(lambda: win.check_updates(silent=False))
        g2, l2 = card("Güncelleme")
        h2 = QHBoxLayout()
        info2 = QVBoxLayout()
        info2.addWidget(self.upd)
        info2.addWidget(self.auto)
        h2.addLayout(info2, 1)
        h2.addWidget(check, 0, Qt.AlignVCenter)
        l2.addLayout(h2)

        # Veriler
        folder = QPushButton("Veri klasörünü aç")
        folder.clicked.connect(lambda: QDesktopServices.openUrl(QUrl.fromLocalFile(data_dir())))
        backup = QPushButton("Yedek al…")
        backup.clicked.connect(self.backup)
        restore = QPushButton("Yedekten geri yükle…")
        restore.clicked.connect(self.restore)
        reset = variant(QPushButton("Tüm verileri sil, sıfırdan başla…"), "danger")
        reset.clicked.connect(self.reset)
        g3, l3 = card("Veriler ve yedek")
        dpath = QLabel("Veriler yalnızca bu bilgisayarda tutulur ve güncellemede silinmez. "
                       "<b>Yedek otomatik alınır:</b> program açılınca, her kayıttan sonra ve "
                       "kapanırken; her gün için bir yedek saklanır (son 30 gün).")
        dpath.setWordWrap(True)
        self.backup_info = QLabel()
        self.backup_info.setObjectName("okMsg")
        self.backup_info.setWordWrap(True)
        open_bk = QPushButton("Yedek klasörünü aç")
        open_bk.clicked.connect(lambda: QDesktopServices.openUrl(
            QUrl.fromLocalFile(self.win.backup_dir())))
        l3.addWidget(dpath)
        l3.addWidget(self.backup_info)
        l3.addLayout(row(open_bk, folder, backup, restore, stretch_first=False))
        l3.addLayout(row(reset, stretch_first=False))

        v = QVBoxLayout(self)
        v.setContentsMargins(0, 0, 0, 0)
        v.setSpacing(12)
        v.addWidget(page_header("Ayarlar"))
        for g in (gt, gm, g1, g4, g2, g3):
            v.addWidget(g)
        v.addStretch()

    def reset(self):
        if not _ask(self, "Ders programı, nöbet çizelgeleri ve TÜM görevlendirme kayıtları "
                          "silinecek. Okul bilgileri ve ayarlar kalacak.\n\nÖnce yedek almak "
                          "isteyebilirsiniz. Silmek istediğinize emin misiniz?", "Sıfırdan başla"):
            return
        self.db.reset_all()
        self.win.sched.refresh()
        self.win.duty.refresh()
        self.win.changed()
        self.win.go_first_step()
        _info(self, "Tüm veriler silindi. Ders Programı sayfasından, ders programını yükleyerek başlayın.")

    def _paint_swatches(self):
        for name, (b, color) in self.swatches.items():
            ring = "3px solid #0f172a" if name == self.accent_name else "2px solid transparent"
            b.setStyleSheet(f"QPushButton {{ background:{color}; border:{ring}; "
                            f"border-radius:19px; padding:0; }}")

    def _pick_accent(self, name):
        self.accent_name = name
        self.db.set_setting("theme_accent", name)
        self._paint_swatches()
        self._theme_changed()

    def _style_choices(self, group):
        for b in group:
            b.setProperty("variant", "primary" if b.isChecked() else "")
            b.style().unpolish(b)
            b.style().polish(b)

    def _choose(self, key, btn, group):
        for b in group:
            b.setChecked(b is btn)
        self._style_choices(group)
        value = btn.text()
        if key == "theme_mode":
            self.db.set_setting("theme_dark", "1" if value == "Koyu" else "0")
        self.db.set_setting(key, value)
        self._theme_changed()

    def _theme_changed(self, *_):
        # Seçilen düğme önce boyanır, tema hemen ardından uygulanır (ardışık tıklamalar
        # tek seferde uygulanır). Sayfalar yeniden kurulmaz, yalnızca stil değişir.
        if not hasattr(self, "_theme_timer"):
            self._theme_timer = QTimer(self)
            self._theme_timer.setSingleShot(True)
            self._theme_timer.timeout.connect(self._apply_theme_now)
        self._theme_timer.start(15)

    def _apply_theme_now(self):
        QApplication.setOverrideCursor(Qt.WaitCursor)
        try:
            self.win.apply_theme()
        finally:
            QApplication.restoreOverrideCursor()

    def open_excel(self):
        path = self.win.excel_path()
        if self.win.write_excel() or os.path.exists(path):
            QDesktopServices.openUrl(QUrl.fromLocalFile(path))

    def move_excel(self):
        folder = QFileDialog.getExistingDirectory(self, "Excel kayıt dosyasının klasörü",
                                                  os.path.dirname(self.win.excel_path()))
        if folder:
            self.db.set_setting("excel_path", os.path.join(folder, export.RECORDS_FILE))
            self.excel.setText(self.win.excel_path())
            self.win.write_excel()

    def backup(self):
        path, _ = QFileDialog.getSaveFileName(
            self, "Yedek kaydet", f"nobet_yedek_{dt.date.today().isoformat()}.db", "Yedek (*.db)")
        if path:
            self.db.conn.commit()
            shutil.copyfile(self.db.path, path)
            _info(self, "Yedek alındı.")

    def restore(self):
        path, _ = QFileDialog.getOpenFileName(self, "Yedek seç", "", "Yedek (*.db)")
        if path and _ask(self, "Mevcut tüm veriler yedektekiyle değiştirilecek. Devam edilsin mi?"):
            self.db.conn.close()
            shutil.copyfile(path, self.db.path)
            self.win.reopen_db()
            _info(self, "Yedek geri yüklendi.")


def notes_html(notes, installed):
    """Sürüm notlarını okunaklı HTML'e çevirir; yalnızca yüklü sürümden yeni olanları alır."""
    import html as _html
    import re as _re

    def ver(text):
        m = _re.search(r"(\d+)\.(\d+)\.(\d+)", text)
        return tuple(int(x) for x in m.groups()) if m else None

    mine = ver(installed) or (0, 0, 0)
    if not any(line.strip().startswith("Sürüm ") and (ver(line) or (0,)) > mine
               for line in (notes or "").splitlines()):
        mine = (-1,)   # yeni bölüm bulunamadı: tümünü göster
    out, keep, in_list = [], True, False
    for raw in (notes or "").splitlines():
        line = raw.strip()
        if line.startswith("Sürüm "):
            v = ver(line)
            keep = v is None or v > mine
            if keep:
                if in_list:
                    out.append("</ul>")
                    in_list = False
                out.append(f"<h3 style='margin:12px 0 4px 0'>🆕 Sürüm {'.'.join(map(str, v))}"
                           "</h3>" if v else f"<h3>{_html.escape(line)}</h3>")
            continue
        if not keep or not line:
            continue
        if line.startswith("- "):
            if not in_list:
                out.append("<ul style='margin:0 0 0 -14px'>")
                in_list = True
            out.append(f"<li style='margin-bottom:5px'>{_html.escape(line[2:])}</li>")
        else:
            if in_list:
                out.append("</ul>")
                in_list = False
            out.append(f"<p style='margin:8px 0 2px 0'><b>{_html.escape(line)}</b></p>")
    if in_list:
        out.append("</ul>")
    return "".join(out) or "<p>Sürüm notu yok.</p>"


class UpdateDialog(QDialog):
    """Yeni sürüm sorusu: yenilikler kaydırılabilir kutuda, Evet/Hayır her zaman görünür."""

    def __init__(self, parent, info):
        super().__init__(parent)
        from PySide6.QtWidgets import QTextBrowser
        self.setWindowTitle("Güncelleme")
        notes = QTextBrowser()
        notes.setOpenExternalLinks(True)
        notes.setHtml("<div style='font-size:10.5pt; line-height:140%'>" +
                      notes_html(info.get("notes", ""), __version__) + "</div>")
        bb = QDialogButtonBox(QDialogButtonBox.Yes | QDialogButtonBox.No)
        bb.button(QDialogButtonBox.Yes).setText("Evet, şimdi güncelle")
        variant(bb.button(QDialogButtonBox.Yes), "primary")
        bb.button(QDialogButtonBox.No).setText("Hayır, sonra")
        bb.accepted.connect(self.accept)
        bb.rejected.connect(self.reject)
        ask = QLabel("Şimdi indirilip kurulsun mu? Program kapanıp yeni sürümle yeniden açılacak; "
                     "verileriniz silinmez.")
        ask.setWordWrap(True)
        v = QVBoxLayout(self)
        v.addWidget(page_header(f"🔄  Yeni sürüm: v{info['version']}",
                                f"Yüklü sürüm: v{__version__}. Yenilikler:"))
        v.addWidget(notes, 1)
        v.addWidget(ask)
        v.addWidget(bb)
        screen = (parent.screen() if parent else QApplication.primaryScreen()).availableGeometry()
        self.resize(min(640, int(screen.width() * 0.9)), min(560, int(screen.height() * 0.8)))


# ---------------------------------------------------------------- ana pencere

class MainWindow(QMainWindow):
    def __init__(self, db=None):
        super().__init__()
        self.setWindowTitle(f"{APP_NAME} — {APP_TAGLINE}  v{__version__}")
        self.resize(1440, 860)
        self.db = db or Database()
        self._workers = []
        if os.path.exists(logo_path()) and self.db.get_setting("logo_round") != "1":
            round_logo(QImage(logo_path())).save(logo_path())  # eski kare logoları yuvarlat
            self.db.set_setting("logo_round", "1")
        self._build()
        QTimer.singleShot(300, self.first_run)
        QTimer.singleShot(800, self.check_new_month)
        QTimer.singleShot(2000, self.auto_backup)
        self._month_timer = QTimer(self)  # program açık kalırsa ay değişimini de yakala
        self._month_timer.timeout.connect(self.check_new_month)
        self._month_timer.start(60 * 60 * 1000)
        if self.db.get_setting("auto_update", "1") == "1" and updater.can_self_update():
            QTimer.singleShot(2500, lambda: self.check_updates(silent=True))

    def _build(self):
        self.tabs = theme.NavPages(APP_NAME, APP_TAGLINE)
        self.tabs.footer.setText(f"Sürüm {__version__}")
        self.school_tab = SchoolTab(self)
        self.daily = DailyTab(self)
        self.sched = ScheduleTab(self)
        self.duty = DutyTab(self)
        self.report = ReportTab(self)
        self.records = RecordsTab(self)
        self.settings = SettingsTab(self)
        # Sıra: önce ders programı, sonra nöbet çizelgesi, sonra görevlendirme
        self.tabs.addTab(self.school_tab, "🏫   Okul Bilgileri")
        self.tabs.addTab(self.sched, "📚   Ders Programı")
        self.tabs.addTab(self.duty, "🗓   Nöbet Çizelgesi")
        self.tabs.addTab(self.daily, "📝   Görevlendirme")
        self.tabs.addTab(self.records, "📋   Kayıtlar")
        self.tabs.addTab(self.report, "📊   Aylık Rapor")
        self.tabs.addTab(self.settings, "⚙   Ayarlar")
        self.tabs.currentChanged.connect(self._tab_changed)
        self.tabs.exit_btn.clicked.connect(self.quit)
        self.setCentralWidget(self.tabs)
        theme.style_root(self.statusBar())
        self.update_brand()
        self.go_first_step()
        self.statusBar().showMessage("Hazır.")

    def first_run(self):
        """İlk açılışta (başka bir okula kurulduğunda) önce okul bilgileri ve logo sorulur."""
        if self.db.get_setting("school"):
            return
        self.tabs.setCurrentWidget(self.school_tab)
        self.statusBar().showMessage(f"👋 {APP_NAME}'e hoş geldiniz! Önce okulunuzun bilgilerini "
                                     "girip Kaydet'e basın.", 20000)

    def school_saved(self, first):
        self.update_brand()
        self.write_excel()
        if first and not self.db.schedule_versions():
            _info(self, "✅ Okul bilgileri kaydedildi.\n\nSıradaki adım: Ders Programı. "
                        "Okulunuzun ders programını PDF veya Excel olarak yükleyin.",
                  "Sonraki adım")
            self.tabs.setCurrentWidget(self.sched)
        else:
            self.statusBar().showMessage("✅ Okul bilgileri kaydedildi.", 6000)

    def quit(self):
        if _ask(self, f"{APP_NAME} kapatılsın mı?", "Çıkış"):
            self.auto_backup()
            QApplication.instance().quit()

    def go_first_step(self):
        """Eksik olan ilk adımı açar: okul → ders programı → nöbet çizelgesi → görevlendirme."""
        if not self.db.get_setting("school"):
            self.tabs.setCurrentWidget(self.school_tab)
        elif not self.db.schedule_versions():
            self.tabs.setCurrentWidget(self.sched)
        elif not self.db.duty_months():
            self.tabs.setCurrentWidget(self.duty)
        else:
            self.tabs.setCurrentWidget(self.daily)

    def ask_admins(self):
        """İdareciler bir kez onaylanmadıysa sorar (seçilmezse onlara da ders verilirdi)."""
        if self.db.get_setting("admins_confirmed") == "1" or not self.db.all_teachers():
            return
        if run_dialog(AdminsDialog(self, self.db)):
            self.sched.refresh_admins()
            self.school_tab.refresh()
            self.daily.refresh()

    def next_step(self, done):
        """Bir adım bitince kullanıcıyı sıradaki adıma götürür."""
        self.ask_admins()
        if not self.db.duty_months():
            _info(self, f"✅ {done}\n\nSıradaki adım: Nöbet Çizelgesi.\n\n"
                        "• otomatik oluşturabilir,\n• hazır çizelgenizi PDF/Excel olarak "
                        "yükleyebilir,\n• ya da elle girebilirsiniz.", "Sonraki adım")
            self.tabs.setCurrentWidget(self.duty)
        else:
            self.statusBar().showMessage(f"✅ {done} Artık görevlendirme yapabilirsiniz.", 8000)
            self.tabs.setCurrentWidget(self.daily)

    def check_new_month(self, force=False):
        """Yeni ay başladıysa (program açıkken de) uyarır ve yeni nöbet çizelgesini önerir."""
        if not force and self.db.get_setting("month_alert", "1") != "1":
            return
        if not self.db.schedule_versions():
            return  # ilk kurulum: zaten adım adım yönlendiriliyor
        month = dt.date.today().strftime("%Y-%m")
        if not force and month in [m for m, _n in self.db.duty_months()]:
            return
        if not force and self.db.get_setting("month_alert_shown") == month and \
                getattr(self, "_alerted", None) == month:
            return
        self._alerted = month
        self.db.set_setting("month_alert_shown", month)
        box = QMessageBox(self)
        box.setIcon(QMessageBox.Information)
        box.setWindowTitle("Yeni ay")
        box.setText(f"<b style='font-size:12pt'>📅 {export.month_tr(month)} başladı.</b><br><br>"
                    "Bu ay için nöbet çizelgesi henüz yok. Ne yapmak istersiniz?<br><br>"
                    "<span style='color:#64748b'>Ders programı değiştiyse önce yeni ders "
                    "programını yükleyin.</span>")
        matik_b = box.addButton("⚙  Otomatik oluştur", QMessageBox.AcceptRole)
        load_b = box.addButton("📄  PDF / Excel'den yükle", QMessageBox.ActionRole)
        edit_b = box.addButton("✏  Elle gir", QMessageBox.ActionRole)
        sched_b = box.addButton("Ders programını güncelle", QMessageBox.ActionRole)
        box.addButton("Sonra", QMessageBox.RejectRole)
        variant(matik_b, "primary")
        box.exec()
        clicked = box.clickedButton()
        box.deleteLater()
        if clicked is matik_b:
            self.duty.open_nobetmatik(month)
        elif clicked is load_b:
            self.tabs.setCurrentWidget(self.duty)
            self.duty.load_file()
        elif clicked is edit_b:
            self.tabs.setCurrentWidget(self.duty)
            self.duty.edit()
        elif clicked is sched_b:
            self.tabs.setCurrentWidget(self.sched)
            self.sched.load_file()

    def apply_theme(self):
        apply_saved_theme(QApplication.instance(), self.db)

    def school(self):
        return self.db.get_setting("school", "")

    def school_info(self):
        return {k: self.db.get_setting(k, "") for k, _l, _h in SCHOOL_KEYS}

    def school_header(self):
        info = self.school_info()
        return " — ".join(x for x in (info["school"], info["school_city"]) if x)

    def edit_school(self):
        self.tabs.setCurrentWidget(self.school_tab)

    def reopen_db(self):
        """Yedek geri yüklendikten sonra veritabanını yeniden açar ve sayfaları yeniler."""
        self.db.__init__(self.db.path)
        self.school_tab.refresh()
        self.sched.refresh()
        self.duty.refresh()
        self.update_brand()
        self.changed()

    def update_brand(self):
        if not hasattr(self, "tabs"):
            return
        name = self.school() or APP_NAME
        pm = QPixmap(logo_path()) if os.path.exists(logo_path()) else theme.app_icon().pixmap(96, 96)
        self.tabs.set_brand(name, f"{APP_NAME} · {APP_TAGLINE}", pm)

    def changed(self):
        self.daily.refresh()
        self.write_excel()

    def backup_dir(self):
        docs = QStandardPaths.writableLocation(QStandardPaths.DocumentsLocation) or \
            os.path.expanduser("~")
        return os.path.join(docs, "Nöbet Görevlendirme", "Yedekler")

    def auto_backup(self):
        """Günün yedeğini alır (aynı gün içinde yenilenir), 30 günden eskileri siler."""
        import glob
        import sqlite3
        folder = self.backup_dir()
        try:
            os.makedirs(folder, exist_ok=True)
            path = os.path.join(folder, f"yedek_{dt.date.today().isoformat()}.db")
            self.db.conn.commit()
            dst = sqlite3.connect(path)
            self.db.conn.backup(dst)
            dst.close()
            for old in sorted(glob.glob(os.path.join(folder, "yedek_*.db")))[:-30]:
                os.remove(old)
        except Exception as e:  # noqa: BLE001 — yedek alınamazsa program çalışmaya devam eder
            self.statusBar().showMessage(f"⚠ Otomatik yedek alınamadı: {e}", 10000)
            return False
        if hasattr(self, "settings"):
            self.settings.backup_info.setText(
                f"✅ Son otomatik yedek: {QDate.currentDate().toString('d MMMM yyyy')} "
                f"{dt.datetime.now().strftime('%H:%M')} — {path}")
        return True

    def closeEvent(self, event):
        self.auto_backup()
        super().closeEvent(event)

    def excel_path(self):
        saved = self.db.get_setting("excel_path")
        if saved:
            return saved
        docs = QStandardPaths.writableLocation(QStandardPaths.DocumentsLocation) or \
            os.path.expanduser("~")
        return os.path.join(docs, "Nöbet Görevlendirme", export.RECORDS_FILE)

    def write_excel(self):
        path = self.excel_path()
        try:
            export.write_records_xlsx(path, self.db, self.school_header())
        except PermissionError:
            self.statusBar().showMessage(
                f"⚠ Excel kayıt dosyası açık olduğu için güncellenemedi: {path} — dosyayı kapatıp "
                "tekrar kaydedin (kayıtlar programda güvende).", 15000)
            return False
        except Exception as e:  # noqa: BLE001
            self.statusBar().showMessage(f"⚠ Excel kayıt dosyası yazılamadı: {e}", 15000)
            return False
        if hasattr(self, "settings"):
            self.settings.excel.setText(path)
        self.auto_backup()
        return True

    def _tab_changed(self, i):
        w = self.tabs.widget(i)
        if w is self.report:
            self.report.refresh()
        elif w is self.records:
            self.records.refresh()
        elif w is self.daily:
            self.daily.refresh()
        elif w is self.duty:
            self.duty._close_matik()
            self.duty.refresh(select=self.duty.month())
        elif w is self.sched:
            self.sched.steps.setCurrentIndex(0)
            self.sched.refresh(select=self.sched.version)
        elif w is self.school_tab:
            self.school_tab.refresh()

    # ---- güncelleme ----
    def _start(self, worker, done, failed=None):
        self._workers.append(worker)
        worker.done.connect(done)
        worker.failed.connect(failed or (lambda m: _warn(self, "Güncelleme", m)))
        worker.finished.connect(lambda: self._workers.remove(worker))
        worker.start()

    def check_updates(self, silent=False):
        def done(info):
            self.settings.upd.setText(f"Yüklü sürüm: <b>{__version__}</b> — GitHub'daki son sürüm: "
                                      f"<b>{info['version'] or '?'}</b>")
            if not updater.is_newer(info["version"]) or not info["url"]:
                if not silent:
                    _info(self, f"Program güncel (v{__version__}).", "Güncelleme")
                return
            if not updater.can_self_update():
                _info(self, f"Yeni sürüm var: v{info['version']}. Kaynak koddan çalıştırıyorsunuz; "
                            "'git pull' ile güncelleyin.", "Güncelleme")
                return
            if run_dialog(UpdateDialog(self, info)):
                self._download(info)

        def failed(msg):
            if not silent:
                self._update_failed(msg)
        self._start(Worker(updater.check), done, failed)

    def _update_failed(self, msg):
        box = QMessageBox(QMessageBox.Warning, "Güncelleme",
                          f"{msg}\n\nYeni sürümü tarayıcıdan da indirebilirsiniz: indirdiğiniz "
                          "NobetGorevlendirme.exe dosyasını eskisinin yerine koyun. "
                          "Verileriniz silinmez.", parent=self)
        open_btn = box.addButton("İndirme sayfasını aç", QMessageBox.AcceptRole)
        box.addButton("Kapat", QMessageBox.RejectRole)
        box.exec()
        if box.clickedButton() is open_btn:
            QDesktopServices.openUrl(QUrl(updater.RELEASES_PAGE))

    def _download(self, info):
        dlg = QProgressDialog("Yeni sürüm indiriliyor…", None, 0, 100, self)
        dlg.setWindowTitle("Güncelleme")
        dlg.setWindowModality(Qt.WindowModal)
        dlg.show()
        w = Worker(updater.download, info, with_progress=True)
        w.progress.connect(lambda d, t: dlg.setValue(int(d * 100 / t) if t else 0))

        def done(path):
            dlg.close()
            self.db.conn.commit()
            updater.apply_and_restart(path)
            QApplication.quit()

        def failed(msg):
            dlg.close()
            self._update_failed(msg)
        self._start(w, done, failed)


def apply_saved_theme(app, db):
    theme.apply(app, db.get_setting("theme_accent", "Mor"), db.get_setting("theme_dark", "0") == "1",
                db.get_setting("theme_font", "Normal"), db.get_setting("theme_density", "Normal"),
                db.get_setting("theme_radius", "Normal"), db.get_setting("theme_family", "Segoe UI"),
                db.get_setting("theme_bg", "Sade"))


def _next_school_day(d):
    while d.weekday() > 4:
        d += dt.timedelta(days=1)
    return d


def _guess_month(entries):
    dates = sorted(e.date for e in entries if e.date)
    if dates:
        return dates[len(dates) // 2][:7]
    return dt.date.today().strftime("%Y-%m")


def run():
    import sys
    if sys.platform == "win32":
        # Görev çubuğu, Python/eski önbellek yerine programın kendi simgesini göstersin
        try:
            import ctypes
            ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID("NobetGorev.Uygulama")
        except Exception:  # noqa: BLE001
            pass
    QLocale.setDefault(QLocale(QLocale.Turkish, QLocale.Turkey))
    app = QApplication(sys.argv)
    # Qt'nin kendi pencereleri (baskı önizleme, dosya seçme, sağ tık menüsü) Türkçe olsun
    from PySide6.QtCore import QLibraryInfo, QTranslator
    tr = QTranslator(app)
    if tr.load(theme.resource("qtbase_tr.qm")) or \
            tr.load("qtbase_tr", QLibraryInfo.path(QLibraryInfo.TranslationsPath)):
        app.installTranslator(tr)
    app.setApplicationName(APP_NAME)
    db = Database()
    apply_saved_theme(app, db)
    win = MainWindow(db)
    win.showMaximized()
    sys.exit(app.exec())
